import streamlit as st
import sqlite3
import hashlib
from datetime import datetime, date, timedelta
import calendar
from io import BytesIO
import base64
import shutil
import secrets
import hmac
import os
import re
from pathlib import Path

try:
    from streamlit_drawable_canvas import st_canvas
except ImportError:
    st_canvas = None

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas as pdf_canvas
    from reportlab.lib.utils import ImageReader
except ImportError:
    A4 = None
    pdf_canvas = None
    ImageReader = None

# ============================================================
# CONFIGURACIÓN GENERAL
# ============================================================

st.set_page_config(
    page_title="TNG - Control de Uniformes",
    page_icon="👕",
    layout="wide"
)

# Identidad visual compacta basada en el diseño aprobado para TNG.
st.markdown("""
<style>
:root { --tng-navy: #092849; --tng-blue: #0878ee; --tng-pale: #f3f8ff; }
[data-testid="stAppViewContainer"] { background: #f4f8fd; }
[data-testid="stHeader"] { background: rgba(244,248,253,.92); }
[data-testid="stSidebar"] { background: var(--tng-navy); }
[data-testid="stSidebar"] * { color: #f4f8ff; }
[data-testid="stSidebar"] [data-testid="stRadio"] label { padding: .45rem .55rem; border-radius: .45rem; }
[data-testid="stSidebar"] [data-testid="stRadio"] label:hover { background: #123d68; }
[data-testid="stSidebar"] hr { border-color: #ffffff30; }
.stButton > button, [data-testid="stFormSubmitButton"] button { background: var(--tng-blue); color: white; border: 0; border-radius: 8px; font-weight: 600; }
.stButton > button:hover, [data-testid="stFormSubmitButton"] button:hover { background: #0566cf; color: white; border: 0; }
[data-testid="stMetric"] { background: white; border: 1px solid #dce7f4; padding: 1rem; border-radius: 12px; box-shadow: 0 3px 12px #102b4808; }
[data-testid="stVerticalBlockBorderWrapper"] { border-radius: 12px; }
[data-testid="stDataFrame"], [data-testid="stTable"] { border: 1px solid #dce7f4; border-radius: 10px; }
h1, h2, h3 { color: #102b4e; }
@media (max-width: 700px) { section.main > div { padding-left: .7rem; padding-right: .7rem; } }
</style>
""", unsafe_allow_html=True)

DB_NAME = os.environ.get("TNG_DB_PATH", str(Path(__file__).resolve().parent / "tng_uniformes.db"))


# ============================================================
# CONEXIÓN A BASE DE DATOS
# ============================================================

def conectar():
    con = sqlite3.connect(DB_NAME, timeout=30)
    con.execute("PRAGMA busy_timeout=30000")
    return con


def hash_password(password):
    return hashlib.sha256(password.encode("utf-8")).hexdigest()


# ============================================================
# CREAR BASE DE DATOS
# ============================================================

def inicializar_bd():
    conexion = conectar()
    cursor = conexion.cursor()

    # Usuarios
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            usuario TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            rol TEXT NOT NULL,
            activo INTEGER DEFAULT 1
        )
    """)

    # Empleados
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS empleados (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            numero_empleado TEXT UNIQUE NOT NULL,
            nombre TEXT NOT NULL,
            puesto TEXT NOT NULL,
            area TEXT NOT NULL,
            talla_camisa TEXT,
            talla_pantalon TEXT,
            fecha_ingreso TEXT,
            activo INTEGER DEFAULT 1
        )
    """)

    # Prendas / inventario
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS prendas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            tipo TEXT NOT NULL,
            talla TEXT NOT NULL,
            existencia INTEGER DEFAULT 0,
            minimo INTEGER DEFAULT 0,
            stock_seguridad INTEGER DEFAULT 0,
            proveedor TEXT,
            activo INTEGER DEFAULT 1,
            UNIQUE(nombre, talla)
        )
    """)

    # Matriz de dotación
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS dotaciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            puesto TEXT NOT NULL,
            prenda_id INTEGER NOT NULL,
            cantidad INTEGER NOT NULL,
            frecuencia TEXT NOT NULL,
            observaciones TEXT,
            FOREIGN KEY(prenda_id) REFERENCES prendas(id)
        )
    """)

    # Entregas
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS entregas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empleado_id INTEGER NOT NULL,
            prenda_id INTEGER NOT NULL,
            cantidad INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            responsable TEXT NOT NULL,
            acuse INTEGER DEFAULT 1,
            observaciones TEXT,
            FOREIGN KEY(empleado_id) REFERENCES empleados(id),
            FOREIGN KEY(prenda_id) REFERENCES prendas(id)
        )
    """)

    # Compatibilidad con bases existentes: agregar campos de firma si faltan.
    columnas_entregas = [fila[1] for fila in cursor.execute(
        "PRAGMA table_info(entregas)"
    ).fetchall()]

    if "firma" not in columnas_entregas:
        cursor.execute("ALTER TABLE entregas ADD COLUMN firma TEXT")

    if "firma_fecha" not in columnas_entregas:
        cursor.execute("ALTER TABLE entregas ADD COLUMN firma_fecha TEXT")

    columnas_empleados = [fila[1] for fila in cursor.execute(
        "PRAGMA table_info(empleados)"
    ).fetchall()]
    if "disciplina" not in columnas_empleados:
        cursor.execute("ALTER TABLE empleados ADD COLUMN disciplina TEXT DEFAULT ''")
    if "tipo_trabajador" not in columnas_empleados:
        # Los registros anteriores quedan sin clasificar; no se supone su régimen laboral.
        cursor.execute("ALTER TABLE empleados ADD COLUMN tipo_trabajador TEXT NOT NULL DEFAULT 'Sin clasificar'")

    columnas_dotaciones = [fila[1] for fila in cursor.execute(
        "PRAGMA table_info(dotaciones)"
    ).fetchall()]
    if "tipo_trabajador" not in columnas_dotaciones:
        # Las reglas anteriores siguen aplicando a todos hasta que se clasifiquen.
        cursor.execute("ALTER TABLE dotaciones ADD COLUMN tipo_trabajador TEXT NOT NULL DEFAULT 'Todos'")

    # Reposiciones
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS reposiciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empleado_id INTEGER NOT NULL,
            prenda_id INTEGER NOT NULL,
            cantidad INTEGER NOT NULL,
            motivo TEXT NOT NULL,
            fecha_solicitud TEXT NOT NULL,
            fecha_entrega TEXT,
            autorizado_por TEXT,
            entregado_por TEXT,
            estado TEXT DEFAULT 'Pendiente',
            observaciones TEXT,
            FOREIGN KEY(empleado_id) REFERENCES empleados(id),
            FOREIGN KEY(prenda_id) REFERENCES prendas(id)
        )
    """)

    # Movimientos / auditoría
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS movimientos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            prenda_id INTEGER NOT NULL,
            tipo TEXT NOT NULL,
            cantidad INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            usuario TEXT NOT NULL,
            descripcion TEXT,
            FOREIGN KEY(prenda_id) REFERENCES prendas(id)
        )
    """)

    # Bitácora de acciones del sistema
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS bitacora (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT NOT NULL,
            usuario TEXT NOT NULL,
            rol TEXT,
            modulo TEXT NOT NULL,
            accion TEXT NOT NULL,
            detalle TEXT
        )
    """)

    # Solicitudes coordinadas: tabla nueva; no altera registros anteriores.
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS solicitudes_almacen (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empleado_id INTEGER NOT NULL,
            prenda_id INTEGER NOT NULL,
            cantidad INTEGER NOT NULL CHECK(cantidad > 0),
            motivo TEXT NOT NULL,
            observaciones TEXT,
            fecha_solicitud TEXT NOT NULL,
            solicitado_por TEXT NOT NULL,
            estado TEXT NOT NULL DEFAULT 'Pendiente',
            revisado_por TEXT,
            fecha_revision TEXT,
            respuesta TEXT,
            entregado_por TEXT,
            fecha_entrega TEXT,
            entrega_id INTEGER UNIQUE,
            FOREIGN KEY(empleado_id) REFERENCES empleados(id),
            FOREIGN KEY(prenda_id) REFERENCES prendas(id),
            FOREIGN KEY(entrega_id) REFERENCES entregas(id)
        )
    """)

    # Migraciones aditivas: las cuentas y solicitudes anteriores se conservan.
    columnas_usuario = [r[1] for r in cursor.execute("PRAGMA table_info(usuarios)").fetchall()]
    if "area" not in columnas_usuario:
        cursor.execute("ALTER TABLE usuarios ADD COLUMN area TEXT DEFAULT ''")
    columnas_sol = [r[1] for r in cursor.execute("PRAGMA table_info(solicitudes_almacen)").fetchall()]
    if "area_solicitante" not in columnas_sol:
        cursor.execute("ALTER TABLE solicitudes_almacen ADD COLUMN area_solicitante TEXT DEFAULT ''")
    if "autorizado_por" not in columnas_sol:
        cursor.execute("ALTER TABLE solicitudes_almacen ADD COLUMN autorizado_por TEXT")
    if "fecha_autorizacion" not in columnas_sol:
        cursor.execute("ALTER TABLE solicitudes_almacen ADD COLUMN fecha_autorizacion TEXT")

    # Nuevas funciones: migraciones aditivas, sin borrar información anterior.
    columnas_sol = {r[1] for r in cursor.execute("PRAGMA table_info(solicitudes_almacen)")}
    if "cantidad_entregada" not in columnas_sol:
        cursor.execute("ALTER TABLE solicitudes_almacen ADD COLUMN cantidad_entregada INTEGER NOT NULL DEFAULT 0")
    cursor.execute("""CREATE TABLE IF NOT EXISTS solicitudes_entregas_parciales (
        id INTEGER PRIMARY KEY AUTOINCREMENT, solicitud_id INTEGER NOT NULL,
        entrega_id INTEGER NOT NULL UNIQUE, cantidad INTEGER NOT NULL CHECK(cantidad > 0),
        fecha TEXT NOT NULL, usuario TEXT NOT NULL,
        FOREIGN KEY(solicitud_id) REFERENCES solicitudes_almacen(id),
        FOREIGN KEY(entrega_id) REFERENCES entregas(id))""")
    # Reconciliar entregas anteriores asociadas a solicitudes; no duplicarlas.
    cursor.execute("""UPDATE solicitudes_almacen SET cantidad_entregada = cantidad
        WHERE estado='Entregada' AND entrega_id IS NOT NULL AND cantidad_entregada=0""")
    cursor.execute("""CREATE TABLE IF NOT EXISTS devoluciones (
        id INTEGER PRIMARY KEY AUTOINCREMENT, empleado_id INTEGER NOT NULL,
        prenda_id INTEGER NOT NULL, cantidad INTEGER NOT NULL CHECK(cantidad > 0),
        fecha TEXT NOT NULL, motivo TEXT NOT NULL, condicion TEXT NOT NULL,
        reintegrada INTEGER NOT NULL DEFAULT 0, usuario TEXT NOT NULL, observaciones TEXT,
        FOREIGN KEY(empleado_id) REFERENCES empleados(id), FOREIGN KEY(prenda_id) REFERENCES prendas(id))""")

    # Ampliaciones aditivas: no se eliminan registros históricos.
    cursor.execute("""CREATE TABLE IF NOT EXISTS cambios_talla (
        id INTEGER PRIMARY KEY AUTOINCREMENT, empleado_id INTEGER NOT NULL,
        campo TEXT NOT NULL, anterior TEXT, nueva TEXT NOT NULL,
        motivo TEXT NOT NULL, fecha TEXT NOT NULL, usuario TEXT NOT NULL,
        FOREIGN KEY(empleado_id) REFERENCES empleados(id))""")
    cursor.execute("""CREATE TABLE IF NOT EXISTS bajas_trabajadores (
        id INTEGER PRIMARY KEY AUTOINCREMENT, empleado_id INTEGER NOT NULL,
        fecha TEXT NOT NULL, motivo TEXT NOT NULL, usuario TEXT NOT NULL,
        FOREIGN KEY(empleado_id) REFERENCES empleados(id))""")
    cursor.execute("""CREATE TABLE IF NOT EXISTS catalogo_motivos_reposicion (
        id INTEGER PRIMARY KEY AUTOINCREMENT, motivo TEXT NOT NULL UNIQUE,
        activo INTEGER NOT NULL DEFAULT 1)""")
    for motivo in ('Desgaste','Daño','Cambio de talla','Cambio de puesto','Pérdida','Otro'):
        cursor.execute('INSERT OR IGNORE INTO catalogo_motivos_reposicion(motivo) VALUES (?)',(motivo,))

    # Migraciones adicionales: recuperación, firmas de autorización y concurrencia.
    cursor.execute("""CREATE TABLE IF NOT EXISTS autorizaciones_firmadas (
        id INTEGER PRIMARY KEY AUTOINCREMENT, solicitud_id INTEGER NOT NULL UNIQUE,
        usuario TEXT NOT NULL, fecha TEXT NOT NULL, firma TEXT NOT NULL,
        comentario TEXT, FOREIGN KEY(solicitud_id) REFERENCES solicitudes_almacen(id))""")
    cursor.execute("""CREATE TABLE IF NOT EXISTS restablecimientos_clave (
        id INTEGER PRIMARY KEY AUTOINCREMENT, usuario TEXT NOT NULL, fecha TEXT NOT NULL,
        autorizado_por TEXT NOT NULL, debe_cambiar INTEGER NOT NULL DEFAULT 1)""")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_movimientos_prenda_fecha ON movimientos(prenda_id,fecha)")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_solicitudes_estado ON solicitudes_almacen(estado)")

    # No crear cuentas con contraseñas públicas. Las cuentas existentes se conservan.
    # Aprovisionar usuarios mediante procedimiento administrativo seguro.
    conexion.commit()
    conexion.close()


# ============================================================
# UTILIDADES
# ============================================================

def ahora():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def registrar_movimiento(prenda_id, tipo, cantidad, usuario, descripcion):
    conexion = conectar()
    cursor = conexion.cursor()

    cursor.execute("""
        INSERT INTO movimientos
        (prenda_id, tipo, cantidad, fecha, usuario, descripcion)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (
        prenda_id,
        tipo,
        cantidad,
        ahora(),
        usuario,
        descripcion
    ))

    conexion.commit()
    conexion.close()


def firma_a_base64(image_data):
    """Convierte el dibujo de la firma a PNG Base64."""
    if image_data is None:
        return None

    try:
        from PIL import Image
        imagen = Image.fromarray(image_data.astype("uint8"), mode="RGBA")
        bbox = imagen.getbbox()
        if bbox:
            imagen = imagen.crop(bbox)
        buffer = BytesIO()
        imagen.save(buffer, format="PNG")
        return base64.b64encode(buffer.getvalue()).decode("utf-8")
    except Exception:
        return None


def generar_pdf_acuse(datos):
    """Genera el acuse de entrega en PDF y devuelve sus bytes."""
    if pdf_canvas is None:
        return None

    buffer = BytesIO()
    pdf = pdf_canvas(buffer, pagesize=A4)
    ancho, alto = A4

    margen = 42
    y = alto - 50

    pdf.setTitle(f"Acuse de entrega ENT-{datos['folio']:06d}")
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(margen, y, "ACUSE DE ENTREGA DE UNIFORME")
    y -= 24
    pdf.setFont("Helvetica", 10)
    pdf.drawString(margen, y, "Sistema de Control de Uniformes")
    pdf.drawRightString(ancho - margen, y, f"Folio: ENT-{datos['folio']:06d}")

    y -= 35
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(margen, y, "Datos del empleado")
    y -= 18
    pdf.setFont("Helvetica", 10)
    campos = [
        ("Matrícula", datos["matricula"]),
        ("Nombre", datos["nombre"]),
        ("Puesto", datos["puesto"]),
        ("Área", datos["area"]),
        ("Disciplina", datos.get("disciplina", "")),
        ("Fecha", datos["fecha"]),
        ("Responsable de entrega", datos["responsable"]),
    ]
    for etiqueta, valor in campos:
        pdf.drawString(margen, y, f"{etiqueta}: {valor}")
        y -= 16

    y -= 8
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(margen, y, "Detalle de la entrega")
    y -= 20
    pdf.setFont("Helvetica", 10)
    pdf.drawString(margen, y, "Prenda")
    pdf.drawString(margen + 230, y, "Talla")
    pdf.drawString(margen + 330, y, "Cantidad")
    y -= 16
    pdf.line(margen, y + 4, ancho - margen, y + 4)
    pdf.drawString(margen, y, str(datos["prenda"]))
    pdf.drawString(margen + 230, y, str(datos["talla"]))
    pdf.drawString(margen + 330, y, str(datos["cantidad"]))

    y -= 35
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(margen, y, "Observaciones")
    y -= 17
    pdf.setFont("Helvetica", 10)
    observaciones = datos.get("observaciones") or "Sin observaciones."
    pdf.drawString(margen, y, observaciones[:100])

    y -= 45
    pdf.setFont("Helvetica", 10)
    texto = "Declaro haber recibido las prendas indicadas anteriormente."
    pdf.drawString(margen, y, texto)

    # Firma
    y -= 28
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(margen, y, "Firma del empleado")
    y -= 10
    if datos.get("firma"):
        try:
            imagen = ImageReader(BytesIO(base64.b64decode(datos["firma"])))
            pdf.drawImage(imagen, margen, y - 85, width=220, height=80,
                          preserveAspectRatio=True, mask='auto')
        except Exception:
            pdf.setFont("Helvetica", 9)
            pdf.drawString(margen, y - 35, "Firma digital registrada")
    else:
        pdf.line(margen, y - 70, margen + 220, y - 70)

    pdf.setFont("Helvetica", 9)
    pdf.drawString(margen, y - 90, f"Empleado: {datos['nombre']}")
    pdf.drawString(margen, y - 104, f"Matrícula: {datos['matricula']}")

    # Responsable
    rx = margen + 285
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(rx, y, "Responsable de entrega")
    pdf.line(rx, y - 70, rx + 190, y - 70)
    pdf.setFont("Helvetica", 9)
    pdf.drawString(rx, y - 90, datos["responsable"][:45])

    pdf.setFont("Helvetica-Oblique", 8)
    pdf.drawString(margen, 35, "Registro generado por el Sistema de Control de Uniformes.")
    pdf.save()
    return buffer.getvalue()


def mostrar_acuse(datos):
    """Muestra el acuse recién generado y el botón para PDF."""
    st.subheader("🧾 Acuse de entrega")
    st.success("✅ La entrega fue registrada y la firma quedó guardada.")
    st.info("♻️ ACUSE DIGITAL: no es necesario imprimir. Puedes conservarlo como registro digital y descargar el PDF solo si necesitas enviarlo o archivarlo.")

    col_izq, col_der = st.columns([1.15, 1])

    with col_izq:
        st.markdown(f"**Folio:** ENT-{datos['folio']:06d}")
        st.markdown(f"**Matrícula:** {datos['matricula']}")
        st.markdown(f"**Empleado:** {datos['nombre']}")
        st.markdown(f"**Puesto:** {datos['puesto']}")
        st.markdown(f"**Área:** {datos['area']}")
        if datos.get("disciplina"):
            st.markdown(f"**Disciplina:** {datos['disciplina']}")
        st.markdown(f"**Prenda:** {datos['prenda']}")
        st.markdown(f"**Talla:** {datos['talla']}")
        st.markdown(f"**Cantidad:** {datos['cantidad']}")
        st.markdown(f"**Fecha:** {datos['fecha']}")
        st.markdown(f"**Responsable:** {datos['responsable']}")
        if datos.get("observaciones"):
            st.markdown(f"**Observaciones:** {datos['observaciones']}")

    with col_der:
        st.markdown("### ✍️ Firma del empleado")
        if datos.get("firma"):
            try:
                from PIL import Image
                imagen = Image.open(BytesIO(base64.b64decode(datos["firma"])))
                st.image(imagen, width=420)
            except Exception:
                st.info("Firma digital registrada.")

    st.divider()
    st.markdown("### 🧾 Confirmación")
    st.write("Declaro haber recibido las prendas indicadas anteriormente.")

    pdf_bytes = generar_pdf_acuse(datos)
    if pdf_bytes:
        st.download_button(
            "📄 Descargar PDF (solo si se necesita)",
            data=pdf_bytes,
            file_name=f"acuse_ENT-{datos['folio']:06d}.pdf",
            mime="application/pdf",
            use_container_width=True
        )
    else:
        st.warning("Instala ReportLab para generar el PDF: python -m pip install reportlab")


def registrar_bitacora(modulo, accion, detalle=""):
    """Registra una acción importante sin interrumpir la operación si falla."""
    try:
        conexion = conectar()
        cursor = conexion.cursor()
        cursor.execute("""
            INSERT INTO bitacora (fecha, usuario, rol, modulo, accion, detalle)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            ahora(),
            st.session_state.get("usuario") or "sistema",
            st.session_state.get("rol") or "",
            modulo,
            accion,
            detalle
        ))
        conexion.commit()
        conexion.close()
    except Exception:
        pass


def cerrar_sesion():
    st.session_state.logged_in = False
    st.session_state.usuario = None
    st.session_state.rol = None
    st.session_state.area = None
    st.rerun()


# ============================================================
# INICIALIZAR
# ============================================================

inicializar_bd()


def respaldo_automatico():
    """Una copia diaria consistente de SQLite; no sobreescribe respaldos anteriores."""
    destino = Path(DB_NAME).resolve().parent / 'respaldos_tng'
    destino.mkdir(exist_ok=True)
    ruta = destino / ('tng_uniformes_' + date.today().isoformat() + '.db')
    if not ruta.exists():
        origen = conectar()
        try:
            with sqlite3.connect(str(ruta)) as copia:
                origen.backup(copia)
        finally:
            origen.close()
    return ruta


def buscar_asistente(pregunta, rol):
    """Consultas de solo lectura, parametrizadas y sin servicios externos."""
    import re
    q = pregunta.strip()
    ql = q.casefold()
    con = conectar()
    try:
        if any(w in ql for w in ('matrícula','matricula','trabajador','empleado','persona','expediente','busca','encuentra')):
            if rol != 'Administrador':
                return 'La búsqueda de datos personales está disponible únicamente para Administrador.'
            patron = re.sub(r'(?i)\b(busca|buscar|encuentra|encontrar|la|el|matrícula|matricula|de|del|trabajador|empleado|persona|expediente|por|nombre|cuál|cual|es|quiero|ver)\b', ' ', q)
            patron = ' '.join(patron.split()).strip(' ?¿')
            if len(patron) < 2:
                return 'Indica al menos dos caracteres del nombre o la matrícula. Ejemplo: Busca Juan Pérez.'
            filas = con.execute('SELECT numero_empleado,nombre,puesto,area,activo FROM empleados WHERE nombre LIKE ? OR numero_empleado LIKE ? ORDER BY activo DESC,nombre LIMIT 10',('%'+patron+'%','%'+patron+'%')).fetchall()
            if not filas:
                return 'No encontré coincidencias. Prueba con parte del nombre, apellido o matrícula.'
            return 'Coincidencias registradas:\n' + '\n'.join(f'- {m} — {n} — {puesto} / {area} — {"Activo" if activo else "Inactivo"}' for m,n,puesto,area,activo in filas)
        if any(w in ql for w in ('stock','existencia','disponible','inventario','hay ','bota','overall','camisola','pantalón','pantalon')):
            terminos = ['duramax','petroleras','mr','overall','camisola','pantalón','pantalon','botas']
            tipo = next((t for t in terminos if t in ql),None)
            if not tipo:
                return 'Indica una prenda. Ejemplo: ¿Hay botas DURAMAX número 27?'
            numeros = re.findall(r'(?<!\d)(?:2[2-9]|3[0-9]|4[0-9]|5[0-2])(?!\d)',ql)
            talla = numeros[-1] if numeros else None
            filas = con.execute('SELECT nombre,talla,existencia FROM prendas WHERE activo=1 AND LOWER(nombre) LIKE ? AND (? IS NULL OR talla=?) ORDER BY nombre,talla LIMIT 30',('%'+tipo+'%',talla,talla)).fetchall()
            if not filas: return 'No hay coincidencias registradas para esa prenda y talla.'
            return 'Existencias registradas:\n'+'\n'.join(f'- {n}, talla {t}: {e} unidades' for n,t,e in filas)
        if any(w in ql for w in ('ayuda','cómo','como','instrucciones')):
            return ("Puedo orientarte: para registrar una entrega utiliza Entregas; para una devolución, "
                    "Devoluciones; para revisar una prenda, Historial por prenda o Kardex. "
                    "Para buscar una matrícula, escribe 'Busca' seguido del nombre (solo Administrador).")
        if any(w in ql for w in ('agotad','bajo stock','stock bajo','abastec')):
            filas=con.execute("""SELECT nombre,talla,existencia,minimo FROM prendas
                WHERE activo=1 AND existencia<=minimo ORDER BY existencia,nombre LIMIT 20""").fetchall()
            return ('Prendas que requieren atención:\\n'+'\\n'.join(
                f'- {n} talla {t}: {e} disponibles; mínimo {m}' for n,t,e,m in filas)) if filas else 'No hay prendas bajo el mínimo registrado.'
        if any(w in ql for w in ('solicitud','pendiente')):
            filas=con.execute("SELECT estado,COUNT(*) FROM solicitudes_almacen GROUP BY estado ORDER BY estado").fetchall()
            return 'Solicitudes por estado:\n'+'\n'.join(f'- {estado}: {n}' for estado,n in filas) if filas else 'No hay solicitudes registradas.'
        return 'Puedo buscar matrículas (solo Administrador), consultar existencias de prendas y tallas o resumir estados de solicitudes. No modifico registros desde el chat.'
    finally:
        con.close()


def excel_historial(df):
    import pandas as pd
    out=BytesIO()
    with pd.ExcelWriter(out,engine="openpyxl") as writer:
        df.to_excel(writer,index=False,sheet_name="Historial")
    return out.getvalue()


def semaforo(existencia, minimo):
    if existencia <= 0:
        return "🔴 Agotado"
    if existencia <= minimo:
        return "🟡 Bajo mínimo"
    return "🟢 Disponible"


def sumar_meses(fecha, meses):
    mes = fecha.month - 1 + meses
    anio = fecha.year + mes // 12
    mes = mes % 12 + 1
    return fecha.replace(year=anio, month=mes, day=min(fecha.day, calendar.monthrange(anio, mes)[1]))


def calcular_proximas_dotaciones(con):
    filas = con.execute("""SELECT e.id,e.numero_empleado,e.nombre,e.puesto,e.tipo_trabajador,
        d.prenda_id,p.nombre,p.talla,d.cantidad,d.frecuencia,e.fecha_ingreso
        FROM empleados e JOIN dotaciones d ON LOWER(TRIM(e.puesto))=LOWER(TRIM(d.puesto))
        JOIN prendas p ON p.id=d.prenda_id
        WHERE e.activo=1 AND p.activo=1 AND (d.tipo_trabajador='Todos' OR d.tipo_trabajador=e.tipo_trabajador)
        ORDER BY e.nombre,p.nombre""").fetchall()
    resultado=[]
    for emp,mat,nombre,puesto,tipo,pid,prenda,talla,cant,frec,ingreso in filas:
        ultima=con.execute("SELECT MAX(substr(fecha,1,10)) FROM entregas WHERE empleado_id=? AND prenda_id=?",(emp,pid)).fetchone()[0]
        referencia=ultima or ingreso
        try: fecha=date.fromisoformat(str(referencia)[:10])
        except (ValueError,TypeError): fecha=None
        if frec == 'Anual': proxima=sumar_meses(fecha,12) if fecha else None
        elif frec == 'Semestral': proxima=sumar_meses(fecha,6) if fecha else None
        elif frec == 'Al ingreso': proxima=(fecha if not ultima else None)
        else: proxima=None
        resultado.append({'Matrícula':mat,'Empleado':nombre,'Puesto':puesto,'Tipo':tipo,
            'Prenda':prenda,'Talla':talla,'Cantidad':cant,'Frecuencia':frec,
            'Última entrega':ultima or 'Sin entrega','Próxima fecha':proxima.isoformat() if proxima else 'Por definir',
            'Estado':('Vencida / pendiente' if proxima and proxima <= date.today() else 'Programada' if proxima else 'Revisión manual')})
    return resultado


def generar_excel_reportes(hojas):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    libro=Workbook(); libro.remove(libro.active)
    for titulo,df in hojas.items():
        ws=libro.create_sheet(titulo[:31]); ws.append(list(df.columns))
        for fila in df.itertuples(index=False,name=None): ws.append(list(fila))
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for celda in ws[1]: celda.font=Font(bold=True,color='FFFFFF'); celda.fill=PatternFill('solid',fgColor='26415C')
        for col in ws.columns:
            ancho=min(48,max(12,max(len(str(c.value or '')) for c in col)+2))
            ws.column_dimensions[get_column_letter(col[0].column)].width=ancho
    salida=BytesIO(); libro.save(salida); return salida.getvalue()


def generar_pdf_reporte(hojas, inicio, fin):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from xml.sax.saxutils import escape
    salida=BytesIO(); doc=SimpleDocTemplate(salida,pagesize=landscape(A4),leftMargin=24,rightMargin=24)
    estilos=getSampleStyleSheet(); contenido=[Paragraph('TNG - Reporte de control de uniformes',estilos['Title']),
        Paragraph(f'Periodo: {escape(str(inicio))} a {escape(str(fin))}',estilos['Normal']),Spacer(1,12)]
    for titulo,df in hojas.items():
        contenido.append(Paragraph(escape(titulo),estilos['Heading2']))
        contenido.append(Paragraph(f'Registros: {len(df)}',estilos['Normal']))
        if not df.empty:
            # Resumen legible: tablas extensas se consultan completas en Excel.
            columnas=list(df.columns)[:7]
            tabla=[[Paragraph(escape(str(c)),estilos['BodyText']) for c in columnas]]
            for fila in df[columnas].head(35).itertuples(index=False,name=None):
                tabla.append([Paragraph(escape(str(v if v is not None else ''))[:100],estilos['BodyText']) for v in fila])
            ancho=(landscape(A4)[0]-48)/len(columnas)
            t=LongTable(tabla,colWidths=[ancho]*len(columnas),repeatRows=1)
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#26415C')),
                ('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),0.3,colors.grey),
                ('VALIGN',(0,0),(-1,-1),'TOP'),('FONTSIZE',(0,0),(-1,-1),7)]))
            contenido.append(t)
            if len(df)>35: contenido.append(Paragraph('Vista previa de 35 filas; el Excel contiene el detalle completo.',estilos['Italic']))
        contenido.append(Spacer(1,16))
    doc.build(contenido); return salida.getvalue()


# ============================================================
# LOGIN
# ============================================================

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False

if "usuario" not in st.session_state:
    st.session_state.usuario = None

if "rol" not in st.session_state:
    st.session_state.rol = None

if "area" not in st.session_state:
    st.session_state.area = None

if "ultimo_acuse" not in st.session_state:
    st.session_state.ultimo_acuse = None


if not st.session_state.logged_in:

    st.title("🏢 TNG")
    st.subheader("Sistema de Control de Uniformes")

    st.write(
        "Control de empleados, inventario, entregas, "
        "reposiciones y trazabilidad."
    )

    st.divider()

    col1, col2, col3 = st.columns([1, 2, 1])

    with col2:

        st.subheader("🔐 Iniciar sesión")

        usuario = st.text_input("Usuario")
        password = st.text_input(
            "Contraseña",
            type="password"
        )

        ingresar = st.button(
            "🔐 Iniciar sesión",
            use_container_width=True
        )

        if ingresar:

            conexion = conectar()
            cursor = conexion.cursor()

            cursor.execute("""
                SELECT usuario, rol, COALESCE(area, '')
                FROM usuarios
                WHERE usuario = ?
                AND password = ?
                AND activo = 1
            """, (
                usuario.strip(),
                hash_password(password)
            ))

            resultado = cursor.fetchone()
            conexion.close()

            if resultado:

                st.session_state.logged_in = True
                st.session_state.usuario = resultado[0]
                st.session_state.rol = resultado[1]
                st.session_state.area = resultado[2]
                registrar_bitacora("Acceso", "Inicio de sesión", "Inicio de sesión correcto")

                st.rerun()

            else:

                st.error(
                    "❌ Usuario o contraseña incorrectos."
                )

    with st.expander("¿Olvidaste tu contraseña?"):
        st.info("Solicita a Administrador un restablecimiento de contraseña. Por seguridad, no se muestra si existe un usuario ni se envían contraseñas por el chatbot.")

    st.caption("Solicita tus credenciales al responsable del sistema. Cambia las claves iniciales antes de publicar la aplicación.")

    st.stop()


try:
    respaldo_automatico()
except (OSError, sqlite3.Error) as error:
    st.warning(f"No fue posible crear el respaldo diario: {error}")

# Obligar a cambiar contraseñas temporales antes de acceder a los módulos.
with conectar() as _con:
    _debe_cambiar=_con.execute("""SELECT COUNT(*) FROM restablecimientos_clave
        WHERE usuario=? AND debe_cambiar=1""",(st.session_state.usuario,)).fetchone()[0]
if _debe_cambiar:
    st.title("🔐 Cambia tu contraseña temporal")
    with st.form("cambio_obligatorio"):
        nueva=st.text_input("Nueva contraseña (mínimo 12 caracteres)",type="password")
        repetida=st.text_input("Confirmar nueva contraseña",type="password")
        guardar=st.form_submit_button("Guardar y continuar")
    if guardar:
        if len(nueva)<12 or nueva!=repetida:
            st.error("Usa al menos 12 caracteres y confirma correctamente.")
        else:
            with conectar() as _con:
                _con.execute("UPDATE usuarios SET password=? WHERE usuario=?",(hash_password(nueva),st.session_state.usuario))
                _con.execute("UPDATE restablecimientos_clave SET debe_cambiar=0 WHERE usuario=?",(st.session_state.usuario,))
            registrar_bitacora("Usuarios","Cambio obligatorio de clave","Contraseña temporal reemplazada")
            st.rerun()
    st.stop()

# ============================================================
# BARRA LATERAL
# ============================================================

st.sidebar.write(
    f"👤 Usuario: **{st.session_state.usuario}**"
)

st.sidebar.write(
    f"🔑 Rol: **{st.session_state.rol}**"
)

st.sidebar.caption("Acceso compartido: todas las áreas deben abrir la misma instancia Streamlit y usar la misma base de datos en el servidor. No compartas el archivo SQLite por una carpeta de red.")
st.sidebar.divider()


# Menú según rol
rol = st.session_state.rol

opciones = []

if rol in ["Administrador"]:
    opciones.append("🔐 Gestión de contraseñas")
    opciones.append("👥 Empleados")
    opciones.extend(["🗂️ Expediente y bajas", "📏 Cambios de talla", "⚙️ Motivos de reposición", "💾 Respaldos"])

if rol in ["Administrador", "Almacén"]:
    opciones.append("📨 Solicitudes a Almacén")

if rol in ["Administrador", "Almacén"]:
    opciones.append("👕 Inventario")

if rol in ["Administrador", "Almacén"]:
    opciones.append("🏭 Almacén")

if rol in ["Administrador", "Almacén"]:
    opciones.append("📦 Entregas")

if rol in ["Administrador", "Almacén"]:
    opciones.append("✍️ Acuses y firmas")

if rol in ["Administrador"]:
    opciones.append("🔄 Reposiciones")

if rol in ["Administrador"]:
    opciones.append("📋 Matriz de Dotación")

if rol in ["Administrador", "Almacén"]:
    opciones.append("📊 Dashboard")

if rol in ["Administrador", "Almacén"]:
    opciones.append("📈 Frecuencias")

if rol in ["Administrador", "Almacén"]:
    opciones.append("📦 Kardex")
    opciones.append("📜 Historial por prenda")

if rol in ["Administrador", "Almacén"]:
    opciones.append("⚠️ Alertas")

if rol in ["Administrador", "Almacén"]:
    opciones.extend(["🚦 Semáforo", "📅 Próximas dotaciones", "📉 Pronóstico", "↩️ Devoluciones"])

if rol in ["Administrador", "Almacén"]:
    opciones.append("📑 Reportes")

if rol in ["Administrador", "Almacén"]:
    opciones.append("📝 Bitácora")

if rol in ["Administrador", "Almacén"]:
    opciones.extend(["🔔 Notificaciones", "📊 Indicadores KPI", "🤖 Asistente TNG 2.0"])

# Cinco secciones principales. Los submódulos se muestran dentro de la página.
GRUPOS_MENU = {
    "🏠 Inicio": [],
    "👥 Trabajadores": ["👥 Empleados", "🗂️ Expediente y bajas", "📏 Cambios de talla", "📋 Matriz de Dotación", "📅 Próximas dotaciones"],
    "📦 Inventario": ["👕 Inventario", "🏭 Almacén", "📦 Kardex", "📜 Historial por prenda", "🚦 Semáforo", "⚠️ Alertas", "📉 Pronóstico", "↩️ Devoluciones"],
    "🚚 Dotaciones": ["📨 Solicitudes a Almacén", "📦 Entregas", "✍️ Acuses y firmas", "🔄 Reposiciones", "⚙️ Motivos de reposición"],
    "⚙️ Administración": ["📊 Dashboard", "📊 Indicadores KPI", "📈 Frecuencias", "📑 Reportes", "📝 Bitácora", "🔔 Notificaciones", "🔐 Gestión de contraseñas", "💾 Respaldos", "🤖 Asistente TNG 2.0"],
}
grupos_visibles = ["🏠 Inicio"] + [
    g for g, modulos in GRUPOS_MENU.items()
    if g != "🏠 Inicio" and any(m in opciones for m in modulos)
]
st.sidebar.markdown("### TNG")
st.sidebar.caption("Control y trazabilidad de uniformes")
if "tng_grupo_destino" in st.session_state:
    destino = st.session_state.pop("tng_grupo_destino")
    if destino in grupos_visibles:
        st.session_state["tng_grupo"] = destino
if st.session_state.get("tng_grupo") not in grupos_visibles:
    st.session_state["tng_grupo"] = "🏠 Inicio"
grupo = st.sidebar.radio("Menú principal", grupos_visibles, key="tng_grupo")
if grupo == "🏠 Inicio":
    opcion = "__inicio__"
else:
    disponibles = [m for m in GRUPOS_MENU[grupo] if m in opciones]
    st.subheader(grupo)
    # Selector compacto en lugar de una fila larga de pestañas.
    modulo_previo = st.session_state.get("tng_modulo_sel_" + grupo)
    indice_previo = disponibles.index(modulo_previo) if modulo_previo in disponibles else 0
    opcion = st.selectbox(
        "¿Qué necesitas hacer?", disponibles, index=indice_previo,
        key="tng_modulo_sel_" + grupo
    )
st.sidebar.divider()

if st.sidebar.button(
    "🚪 Cerrar sesión",
    use_container_width=True
):
    cerrar_sesion()


# ============================================================
# EMPLEADOS
# ============================================================

if opcion == "__inicio__":
    st.title("TNG | Control y Trazabilidad de Uniformes")
    st.caption("Bienvenido. Selecciona una sección del menú o utiliza un acceso rápido.")
    if rol in ("Administrador", "Almacén"):
        with conectar() as con:
            total_trabajadores = con.execute("SELECT COUNT(*) FROM empleados").fetchone()[0]
            total_prendas = con.execute("SELECT COUNT(*) FROM prendas").fetchone()[0]
        c1, c2 = st.columns(2)
        c1.metric("Trabajadores registrados", total_trabajadores)
        c2.metric("Referencias de prendas", total_prendas)
    st.subheader("Accesos rápidos")
    cols = st.columns(3)
    accesos = [g for g in grupos_visibles if g != "🏠 Inicio"]
    for i, g in enumerate(accesos):
        with cols[i % 3]:
            if st.button(g, use_container_width=True, key="inicio_" + g):
                st.session_state["tng_grupo_destino"] = g
                st.rerun()
    st.info("Solo aparecen las secciones autorizadas para tu rol.")

elif opcion == "🔐 Gestión de contraseñas":
    st.title("🔐 Recuperación administrada de contraseñas")
    st.caption("Solo Administrador puede generar una contraseña temporal. El usuario deberá cambiarla al iniciar sesión.")
    with conectar() as con:
        cuentas=con.execute("SELECT usuario,rol FROM usuarios WHERE activo=1 ORDER BY usuario").fetchall()
    cuenta=st.selectbox("Cuenta a restablecer",[u for u,r in cuentas])
    confirmacion=st.checkbox("He verificado la identidad de la persona por un procedimiento interno")
    if st.button("Generar contraseña temporal",disabled=not confirmacion):
        temporal=secrets.token_urlsafe(12)
        with conectar() as con:
            con.execute("BEGIN IMMEDIATE")
            con.execute("UPDATE usuarios SET password=? WHERE usuario=? AND activo=1",(hash_password(temporal),cuenta))
            con.execute("INSERT INTO restablecimientos_clave(usuario,fecha,autorizado_por) VALUES (?,?,?)",
                        (cuenta,ahora(),st.session_state.usuario))
        registrar_bitacora("Usuarios","Restablecimiento de clave",f"Cuenta: {cuenta}")
        st.warning("Copia la contraseña temporal ahora; no se guardará en texto legible.")
        st.code(temporal)
    st.divider()
    st.subheader("Cambiar mi contraseña")
    with st.form("cambiar_clave_propia"):
        actual=st.text_input("Contraseña actual",type="password")
        nueva=st.text_input("Nueva contraseña (mínimo 12 caracteres)",type="password")
        confirmar=st.text_input("Repite la nueva contraseña",type="password")
        cambiar=st.form_submit_button("Actualizar contraseña")
    if cambiar:
        with conectar() as con:
            row=con.execute("SELECT password FROM usuarios WHERE usuario=?",(st.session_state.usuario,)).fetchone()
            if not row or not hmac.compare_digest(row[0],hash_password(actual)):
                st.error("La contraseña actual no coincide.")
            elif len(nueva)<12 or nueva!=confirmar:
                st.error("La nueva contraseña debe tener al menos 12 caracteres y coincidir.")
            else:
                con.execute("UPDATE usuarios SET password=? WHERE usuario=?",(hash_password(nueva),st.session_state.usuario))
                con.execute("UPDATE restablecimientos_clave SET debe_cambiar=0 WHERE usuario=?",(st.session_state.usuario,))
                st.success("Contraseña actualizada.")
                registrar_bitacora("Usuarios","Cambio de contraseña","Cuenta propia")

elif opcion == "📜 Historial por prenda":
    st.title("📜 Historial de movimientos por prenda")
    with conectar() as con:
        prendas_hist=con.execute("SELECT id,nombre,talla,existencia FROM prendas ORDER BY nombre,talla").fetchall()
    if prendas_hist:
        mapa={f"{n} · Talla {t} · ID {pid}":(pid,n,t,e) for pid,n,t,e in prendas_hist}
        sel=st.selectbox("Prenda y talla",list(mapa))
        pid,n,t,e=mapa[sel]
        with conectar() as con:
            filas=con.execute("""SELECT fecha,tipo,cantidad,usuario,COALESCE(descripcion,'')
                FROM movimientos WHERE prenda_id=? ORDER BY fecha DESC,id DESC""",(pid,)).fetchall()
        st.metric("Existencia actual",e)
        import pandas as pd
        df=pd.DataFrame(filas,columns=["Fecha","Tipo","Cantidad","Usuario","Descripción"])
        st.dataframe(df,use_container_width=True,hide_index=True)
        st.download_button("Descargar historial Excel",data=excel_historial(df),
            file_name=f"historial_prenda_{pid}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else: st.info("No hay prendas registradas.")

elif opcion == "👥 Empleados":

    st.title("👥 Gestión de Empleados")

    pestaña1, pestaña2, pestaña3 = st.tabs([
        "➕ Registrar empleado",
        "📋 Consultar empleados",
        "🏷️ Clasificar trabajadores existentes"
    ])

    with pestaña1:

        st.subheader("Registrar nuevo empleado")

        with st.form("form_empleado"):

            col1, col2 = st.columns(2)

            with col1:

                numero = st.text_input(
                    "Matrícula *"
                )

                nombre = st.text_input(
                    "Nombre completo *"
                )

                puesto = st.text_input(
                    "Puesto *"
                )

                area = st.text_input(
                    "Área *"
                )

                disciplina = st.text_input(
                    "Disciplina"
                )

                tipo_trabajador = st.selectbox(
                    "Tipo de trabajador *",
                    ["Selecciona un tipo", "Sindicalizado", "Base"],
                    help="Selecciona el régimen laboral que corresponda según los registros autorizados de TNG."
                )

            with col2:

                talla_camisa = st.selectbox(
                    "Talla de camisola",
                    ["", "28", "30", "32", "34", "36", "38", "40",
                     "42", "44", "46", "48", "50", "52"]
                )

                talla_pantalon = st.selectbox(
                    "Talla de pantalón",
                    ["", "28", "30", "32", "34", "36", "38", "40",
                     "42", "44", "46", "48", "50", "52"]
                )

                fecha_ingreso = st.date_input(
                    "Fecha de ingreso"
                )

            guardar = st.form_submit_button(
                "💾 Guardar empleado",
                use_container_width=True
            )

            if guardar:

                if not numero.strip() or not nombre.strip() \
                        or not puesto.strip() or not area.strip() \
                        or tipo_trabajador == "Selecciona un tipo":

                    st.error(
                        "❌ Completa todos los campos obligatorios."
                    )

                else:

                    conexion = conectar()
                    cursor = conexion.cursor()

                    try:

                        cursor.execute("""
                            INSERT INTO empleados
                            (numero_empleado, nombre, puesto, area, disciplina,
                             talla_camisa, talla_pantalon, fecha_ingreso, tipo_trabajador)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (
                            numero.strip(),
                            nombre.strip(),
                            puesto.strip(),
                            area.strip(),
                            disciplina.strip(),
                            talla_camisa,
                            talla_pantalon,
                            str(fecha_ingreso),
                            tipo_trabajador
                        ))

                        conexion.commit()
                        registrar_bitacora("Empleados", "Alta de trabajador", f"Matrícula: {numero.strip()} | Nombre: {nombre.strip()}")

                        st.success(
                            "✅ Empleado registrado correctamente."
                        )

                    except sqlite3.IntegrityError:

                        st.error(
                            "❌ La matrícula ya existe."
                        )

                    finally:

                        conexion.close()

    with pestaña2:

        st.subheader("📋 Empleados registrados")

        conexion = conectar()
        cursor = conexion.cursor()

        cursor.execute("""
            SELECT
                numero_empleado,
                nombre,
                puesto,
                area,
                disciplina,
                tipo_trabajador,
                talla_camisa,
                talla_pantalon,
                fecha_ingreso
            FROM empleados
            WHERE activo = 1
            ORDER BY nombre
        """)

        empleados = cursor.fetchall()
        conexion.close()

        if empleados:

            import pandas as pd

            df = pd.DataFrame(
                empleados,
                columns=[
                    "Matrícula",
                    "Nombre",
                    "Puesto",
                    "Área",
                    "Disciplina",
                    "Tipo de trabajador",
                    "Talla camisola",
                    "Talla pantalón",
                    "Fecha ingreso"
                ]
            )

            st.dataframe(
                df,
                use_container_width=True,
                hide_index=True
            )

        else:

            st.info("No hay empleados registrados.")

    with pestaña3:
        st.subheader("Clasificar trabajadores ya registrados")
        st.caption("Los registros anteriores se conservan. Puedes asignarles su tipo de trabajador sin volver a capturarlos.")
        conexion = conectar()
        trabajadores_clasificar = conexion.execute(
            "SELECT id, numero_empleado, nombre, COALESCE(tipo_trabajador, 'Sin clasificar') "
            "FROM empleados WHERE activo = 1 ORDER BY nombre"
        ).fetchall()
        conexion.close()
        if not trabajadores_clasificar:
            st.info("No hay trabajadores registrados para clasificar.")
        else:
            trabajadores_dict = {
                f"{t[1]} - {t[2]} (actual: {t[3]})": t for t in trabajadores_clasificar
            }
            with st.form("form_clasificar_trabajador"):
                trabajador_key = st.selectbox("Selecciona al trabajador", list(trabajadores_dict.keys()))
                tipo_nuevo = st.selectbox("Tipo de trabajador", ["Sindicalizado", "Base"])
                guardar_tipo = st.form_submit_button("💾 Guardar clasificación", use_container_width=True)
            if guardar_tipo:
                trabajador = trabajadores_dict[trabajador_key]
                conexion = conectar()
                try:
                    conexion.execute(
                        "UPDATE empleados SET tipo_trabajador = ? WHERE id = ?",
                        (tipo_nuevo, trabajador[0])
                    )
                    conexion.commit()
                    registrar_bitacora("Empleados", "Clasificación laboral actualizada", f"Matrícula: {trabajador[1]} | Tipo: {tipo_nuevo}")
                    st.success(f"Se actualizó {trabajador[1]} como {tipo_nuevo}.")
                    st.rerun()
                except Exception as exc:
                    conexion.rollback()
                    st.error(f"No se pudo actualizar la clasificación: {exc}")
                finally:
                    conexion.close()


# ============================================================
# SOLICITUDES ENTRE RELACIONES LABORALES Y ALMACÉN
# ============================================================
elif opcion == "📨 Solicitudes a Almacén":
    st.title("📨 Solicitudes de uniformes")
    st.caption("Las áreas solicitan; Administrador autoriza; Almacén prepara y entrega con firma. El inventario se descuenta únicamente al confirmar la entrega.")
    conexion = conectar()
    empleados_sol = conexion.execute("SELECT id, numero_empleado, nombre, puesto, area, COALESCE(disciplina, ''), COALESCE(tipo_trabajador, 'Sin clasificar') FROM empleados WHERE activo=1 ORDER BY nombre").fetchall()
    prendas_sol = conexion.execute("SELECT id, nombre, talla, existencia FROM prendas WHERE activo=1 ORDER BY nombre, talla").fetchall()
    conexion.close()

    if False:  # Solo dos roles habilitados: Administrador y Almacén. Se conserva el historial de solicitudes.
        with st.expander("👤 Crear acceso para otra área"):
            with st.form("crear_usuario_area"):
                nuevo_usuario = st.text_input("Usuario de área *").strip()
                nueva_clave = st.text_input("Contraseña inicial *", type="password")
                nueva_area = st.text_input("Nombre del área *").strip()
                crear_area = st.form_submit_button("Crear cuenta de solicitante")
            if crear_area:
                if not nuevo_usuario or len(nueva_clave) < 8 or not nueva_area:
                    st.error("Completa usuario y área; la contraseña debe tener al menos 8 caracteres.")
                else:
                    con = conectar()
                    try:
                        con.execute("INSERT INTO usuarios (usuario,password,rol,area) VALUES (?,?,?,?)",
                            (nuevo_usuario,hash_password(nueva_clave),"Solicitante de área",nueva_area))
                        con.commit()
                        registrar_bitacora("Usuarios", "Alta de solicitante", f"Área: {nueva_area} | Usuario: {nuevo_usuario}")
                        st.success("Cuenta de área creada. Comparte las credenciales por un medio privado.")
                    except sqlite3.IntegrityError:
                        st.error("Ese nombre de usuario ya existe.")
                    finally:
                        con.close()

    if rol == "Administrador":
        st.subheader("Nueva solicitud")
        if not empleados_sol or not prendas_sol:
            st.warning("Registra primero trabajadores y prendas en sus respectivos módulos.")
        else:
            if rol == "Solicitante de área":
                empleados_elegibles = [e for e in empleados_sol if e[4].strip().casefold() == (st.session_state.area or "").strip().casefold()]
            else:
                empleados_elegibles = empleados_sol
            emps = {f"{e[1]} - {e[2]} ({e[6]})": e for e in empleados_elegibles}
            if not emps:
                st.info("No hay trabajadores registrados en tu área. Solicita a Administrador actualizar el área de los empleados.")
            prendas_sel = {f"{p[1]} | Talla {p[2]} | Existencia {p[3]}": p for p in prendas_sol}
            if emps:
              emp_key = st.selectbox("Trabajador (matrícula)", list(emps), key="solicitud_empleado_tipo")
              empleado_solicitud = emps[emp_key]
              st.caption(f"Tipo de trabajador: {empleado_solicitud[6]} · Puesto: {empleado_solicitud[3]}")
              con_reglas = conectar()
              reglas_aplicables = con_reglas.execute("""SELECT p.nombre, p.talla, d.cantidad, d.frecuencia, COALESCE(d.tipo_trabajador, 'Todos'), COALESCE(d.observaciones, '')
                  FROM dotaciones d JOIN prendas p ON p.id=d.prenda_id
                  WHERE lower(trim(d.puesto))=lower(trim(?))
                    AND (COALESCE(d.tipo_trabajador, 'Todos')='Todos' OR COALESCE(d.tipo_trabajador, 'Todos')=?)
                  ORDER BY p.nombre, p.talla""", (empleado_solicitud[3], empleado_solicitud[6])).fetchall()
              con_reglas.close()
              with st.expander("📋 Consultar matriz de dotación aplicable", expanded=True):
                  if reglas_aplicables:
                      import pandas as pd
                      st.dataframe(pd.DataFrame(reglas_aplicables, columns=["Prenda", "Talla", "Cantidad autorizada", "Frecuencia", "Tipo de trabajador", "Observaciones"]), use_container_width=True, hide_index=True)
                  else:
                      st.info("No hay reglas registradas para este puesto y tipo de trabajador. Administrador debe revisar la matriz antes de autorizar la solicitud.")
              with st.form("solicitud_rl"):
                  pre_key = st.selectbox("Prenda solicitada", list(prendas_sel))
                  cant = st.number_input("Cantidad solicitada", min_value=1, step=1, value=1)
                  motivo = st.selectbox("Motivo", ["Dotación inicial", "Dotación programada", "Cambio de talla", "Reposición", "Otro"])
                  obs = st.text_area("Observaciones")
                  enviar = st.form_submit_button("📨 Enviar solicitud", use_container_width=True)
              if enviar:
                  con = conectar()
                  try:
                      con.execute("""INSERT INTO solicitudes_almacen
                          (empleado_id,prenda_id,cantidad,motivo,observaciones,fecha_solicitud,solicitado_por,area_solicitante,estado)
                          VALUES (?,?,?,?,?,?,?,?,?)""", (emps[emp_key][0], prendas_sel[pre_key][0], int(cant), motivo, obs.strip(), ahora(), st.session_state.usuario, st.session_state.area if rol == "Solicitante de área" else emps[emp_key][4], "Por autorizar" if rol == "Solicitante de área" else "Pendiente"))
                      folio = con.execute("SELECT last_insert_rowid()").fetchone()[0]
                      con.commit()
                      registrar_bitacora("Solicitudes a Almacén", "Nueva solicitud", f"Solicitud #{folio}")
                      st.success(f"Solicitud #{folio} registrada. Estado: {'Por autorizar' if rol == 'Solicitante de área' else 'Pendiente'}.")
                  except Exception as exc:
                      con.rollback(); st.error(f"No se pudo registrar: {exc}")
                  finally:
                      con.close()
    con = conectar()
    solicitudes = con.execute("""SELECT s.id, e.numero_empleado, e.nombre, e.puesto, e.area, COALESCE(e.disciplina,''),
        p.nombre, p.talla, s.cantidad, s.motivo, COALESCE(s.observaciones,''), s.fecha_solicitud,
        s.solicitado_por, s.estado, COALESCE(s.respuesta,''), p.existencia, s.empleado_id, s.prenda_id,
        COALESCE(s.entregado_por,''), COALESCE(s.fecha_entrega,''), s.entrega_id, COALESCE(s.area_solicitante,'')
        FROM solicitudes_almacen s JOIN empleados e ON e.id=s.empleado_id JOIN prendas p ON p.id=s.prenda_id
        ORDER BY s.id DESC""").fetchall()
    con.close()
    if rol == "Solicitante de área":
        solicitudes = [x for x in solicitudes if x[12] == st.session_state.usuario and x[21].strip().casefold() == (st.session_state.area or "").strip().casefold()]
    st.divider()
    st.subheader("Seguimiento de solicitudes")
    filtro_sol = st.selectbox("Estado", ["Todas", "Por autorizar", "Pendiente", "Preparada", "Parcial", "Entregada", "Rechazada"], key="filtro_sol")
    visibles = [x for x in solicitudes if filtro_sol == "Todas" or x[13] == filtro_sol]
    if not visibles:
        st.info("No hay solicitudes para este filtro.")
    for x in visibles:
        (sid, matricula, nombre_emp, puesto_emp, area_emp, disciplina_emp, nombre_pre, talla_pre,
         cantidad_sol, motivo_sol, obs_sol, fecha_sol, autor_sol, estado_sol, respuesta_sol, stock_sol,
         emp_id, pre_id, entregado_sol, fecha_ent_sol, entrega_id_sol, area_solicitante) = x
        with st.expander(f"#{sid} · {estado_sol} · {matricula} - {nombre_emp} · {nombre_pre} ({talla_pre})"):
            st.write(f"**Tipo de trabajador:** {next((e[6] for e in empleados_sol if e[1] == matricula), 'Sin clasificar')}")
            st.write(f"**Solicitado:** {cantidad_sol} | **Motivo:** {motivo_sol} | **Fecha:** {fecha_sol} | **Por:** {autor_sol}")
            if obs_sol: st.write(f"**Observaciones:** {obs_sol}")
            if respuesta_sol: st.write(f"**Respuesta de Almacén:** {respuesta_sol}")
            if fecha_ent_sol: st.write(f"**Entregado:** {fecha_ent_sol} | **Responsable:** {entregado_sol} | **Entrega #{entrega_id_sol}")
            con_hist=conectar()
            parcial_hist=con_hist.execute("SELECT cantidad_entregada FROM solicitudes_almacen WHERE id=?",(sid,)).fetchone()[0]
            historial=con_hist.execute("SELECT entrega_id,cantidad,fecha,usuario FROM solicitudes_entregas_parciales WHERE solicitud_id=? ORDER BY id",(sid,)).fetchall()
            con_hist.close()
            st.write(f"**Avance:** {parcial_hist} de {cantidad_sol} unidades | **Pendiente:** {max(0,cantidad_sol-parcial_hist)}")
            if historial:
                st.dataframe([{'Acuse / entrega':h[0],'Cantidad':h[1],'Fecha':h[2],'Responsable':h[3]} for h in historial],hide_index=True,use_container_width=True)
            if rol == "Administrador" and estado_sol == "Por autorizar":
                st.caption("Para autorizar se requiere firma manuscrita del responsable. El rechazo no requiere firma.")
                firma_aut_canvas = None
                if st_canvas is not None:
                    firma_aut_canvas = st_canvas(fill_color="rgba(255,255,255,0)", stroke_width=3,
                        stroke_color="#111111", background_color="#FFFFFF", height=160, width=540,
                        drawing_mode="freedraw", display_toolbar=True, key=f"firma_autorizacion_{sid}")
                else:
                    st.warning("Para autorizar instala streamlit-drawable-canvas y numpy.")
                with st.form(f"autorizacion_area_{sid}"):
                    comentario_rl = st.text_area("Comentario de Administrador")
                    ca, cr = st.columns(2)
                    autorizar_area = ca.form_submit_button("✅ Autorizar y enviar a Almacén")
                    rechazar_area = cr.form_submit_button("❌ Rechazar solicitud")
                if autorizar_area or rechazar_area:
                    import numpy as np
                    pix = firma_aut_canvas.image_data if firma_aut_canvas is not None else None
                    firma_valida = pix is not None and bool(np.any(np.asarray(pix)[:, :, :3] < 180))
                    if rechazar_area and not comentario_rl.strip():
                        st.error("Escribe el motivo del rechazo.")
                    elif autorizar_area and not firma_valida:
                        st.error("Dibuja tu firma antes de autorizar.")
                    else:
                        con = conectar()
                        try:
                            estado_nuevo = "Pendiente" if autorizar_area else "Rechazada"
                            cur = con.execute("""UPDATE solicitudes_almacen SET estado=?, autorizado_por=?, fecha_autorizacion=?, respuesta=?
                                WHERE id=? AND estado='Por autorizar'""", (estado_nuevo,st.session_state.usuario,ahora(),comentario_rl.strip(),sid))
                            if cur.rowcount and autorizar_area:
                                con.execute("""INSERT INTO autorizaciones_firmadas(solicitud_id,usuario,fecha,firma,comentario)
                                    VALUES (?,?,?,?,?)""",(sid,st.session_state.usuario,ahora(),
                                    firma_a_base64(pix),comentario_rl.strip()))
                            con.commit()
                            if cur.rowcount:
                                registrar_bitacora("Solicitudes a Almacén", estado_nuevo, f"Solicitud de área #{sid}")
                                st.success("Solicitud actualizada."); st.rerun()
                            else: st.warning("La solicitud ya fue procesada.")
                        except Exception as exc:
                            con.rollback(); st.error(str(exc))
                        finally: con.close()
            if rol == "Almacén" and estado_sol == "Pendiente":
                st.write(f"Existencia actual: {stock_sol}")
                with st.form(f"revision_sol_{sid}"):
                    respuesta = st.text_area("Comentario de revisión", key=f"resp_{sid}")
                    c1,c2 = st.columns(2)
                    preparar = c1.form_submit_button("✅ Marcar preparada", use_container_width=True)
                    rechazar = c2.form_submit_button("❌ Rechazar", use_container_width=True)
                if preparar or rechazar:
                    if preparar and stock_sol <= 0:
                        st.error("No hay existencia disponible. Registra una entrada antes de preparar la solicitud.")
                    elif rechazar and not respuesta.strip():
                        st.error("Indica el motivo del rechazo.")
                    else:
                        nuevo_estado = "Preparada" if preparar else "Rechazada"
                        con = conectar()
                        try:
                            cur = con.execute("""UPDATE solicitudes_almacen SET estado=?, revisado_por=?, fecha_revision=?, respuesta=?
                                WHERE id=? AND estado='Pendiente'""", (nuevo_estado, st.session_state.usuario, ahora(), respuesta.strip(), sid))
                            con.commit()
                            if cur.rowcount:
                                registrar_bitacora("Solicitudes a Almacén", nuevo_estado, f"Solicitud #{sid}")
                                st.success(f"Solicitud #{sid}: {nuevo_estado}."); st.rerun()
                            else: st.warning("Esta solicitud ya fue revisada. Actualiza la página.")
                        except Exception as exc:
                            con.rollback(); st.error(str(exc))
                        finally: con.close()
            if rol == "Almacén" and estado_sol in ("Preparada", "Parcial"):
                con_tmp=conectar()
                entregada_tmp=con_tmp.execute("SELECT cantidad_entregada FROM solicitudes_almacen WHERE id=?",(sid,)).fetchone()[0]
                con_tmp.close()
                restante_tmp=max(0,cantidad_sol-entregada_tmp)
                st.write(f"Solicitado: {cantidad_sol} | Entregado: {entregada_tmp} | Pendiente: {restante_tmp}")
                cantidad_parcial=st.number_input("Cantidad a entregar ahora",min_value=1,max_value=max(1,restante_tmp),value=max(1,min(restante_tmp,stock_sol)),key=f"parcial_{sid}")
                st.write("El trabajador debe firmar cada entrega, incluso si es parcial.")
                if st_canvas is None:
                    st.warning("Instala streamlit-drawable-canvas y pillow para habilitar la firma digital.")
                else:
                    canvas_sol = st_canvas(fill_color="rgba(255,255,255,0)", stroke_width=3,
                        stroke_color="#111111", background_color="#FFFFFF", height=180, width=600,
                        drawing_mode="freedraw", display_toolbar=True, key=f"firma_solicitud_{sid}")
                    if st.button("📦 Confirmar entrega y descontar inventario", key=f"entregar_sol_{sid}"):
                        firma_sol = firma_a_base64(canvas_sol.image_data if canvas_sol else None)
                        # Comprobar trazos reales; un fondo blanco no equivale a una firma.
                        import numpy as np
                        pixels = canvas_sol.image_data if canvas_sol else None
                        tiene_trazos = pixels is not None and np.any(np.asarray(pixels)[:, :, :3] < 180)
                        if not firma_sol or not tiene_trazos:
                            st.error("El trabajador debe firmar antes de confirmar.")
                        else:
                            con = conectar()
                            try:
                                con.execute("BEGIN IMMEDIATE")
                                actual = con.execute("SELECT estado, empleado_id, prenda_id, cantidad FROM solicitudes_almacen WHERE id=?", (sid,)).fetchone()
                                if not actual or actual[0] not in ("Preparada", "Parcial"):
                                    raise ValueError("La solicitud ya fue procesada o cambió de estado.")
                                entregado_previo=con.execute("SELECT cantidad_entregada FROM solicitudes_almacen WHERE id=?",(sid,)).fetchone()[0]
                                if cantidad_parcial > actual[3]-entregado_previo:
                                    raise ValueError("La cantidad supera el saldo pendiente.")
                                cur = con.execute("UPDATE prendas SET existencia=existencia-? WHERE id=? AND activo=1 AND existencia>=?", (cantidad_parcial, actual[2], cantidad_parcial))
                                if cur.rowcount != 1:
                                    raise ValueError("Inventario insuficiente. No se registró la entrega.")
                                fecha_hora = ahora()
                                cur = con.execute("""INSERT INTO entregas
                                    (empleado_id,prenda_id,cantidad,fecha,responsable,acuse,observaciones,firma,firma_fecha)
                                    VALUES (?,?,?,?,?,1,?,?,?)""", (actual[1],actual[2],cantidad_parcial,fecha_hora,st.session_state.usuario,
                                    f"Solicitud de Almacén #{sid}. {obs_sol}",firma_sol,fecha_hora))
                                folio_ent = cur.lastrowid
                                con.execute("""INSERT INTO movimientos (prenda_id,tipo,cantidad,fecha,usuario,descripcion)
                                    VALUES (?, 'SALIDA', ?, ?, ?, ?)""", (actual[2],cantidad_parcial,fecha_hora,st.session_state.usuario,f"Entrega por solicitud #{sid}; entrega #{folio_ent}"))
                                nuevo_total=entregado_previo+cantidad_parcial
                                nuevo_estado='Entregada' if nuevo_total==actual[3] else 'Parcial'
                                con.execute("""INSERT INTO solicitudes_entregas_parciales
                                    (solicitud_id,entrega_id,cantidad,fecha,usuario) VALUES (?,?,?,?,?)""",
                                    (sid,folio_ent,cantidad_parcial,fecha_hora,st.session_state.usuario))
                                con.execute("""UPDATE solicitudes_almacen SET estado=?, cantidad_entregada=?, entregado_por=?, fecha_entrega=?, entrega_id=?
                                    WHERE id=? AND estado IN ('Preparada','Parcial')""", (nuevo_estado,nuevo_total,st.session_state.usuario,fecha_hora,folio_ent,sid))
                                con.commit()
                                st.session_state.ultimo_acuse = {"folio":folio_ent,"matricula":matricula,"nombre":nombre_emp,
                                    "puesto":puesto_emp,"area":area_emp,"disciplina":disciplina_emp,"prenda":nombre_pre,"talla":talla_pre,
                                    "cantidad":cantidad_parcial,"fecha":fecha_hora,"responsable":st.session_state.usuario,
                                    "observaciones":f"Solicitud #{sid}. {obs_sol}","firma":firma_sol}
                                registrar_bitacora("Solicitudes a Almacén", "Entrega confirmada", f"Solicitud #{sid} | Entrega #{folio_ent}")
                                st.success(f"Entrega #{folio_ent} registrada correctamente."); st.rerun()
                            except Exception as exc:
                                con.rollback(); st.error(f"No se realizó ningún descuento: {exc}")
                            finally: con.close()
            if estado_sol == "Entregada" and entrega_id_sol:
                if st.button("🧾 Ver acuse digital", key=f"acuse_sol_{sid}"):
                    con = conectar()
                    entrega = con.execute("SELECT fecha,responsable,observaciones,firma FROM entregas WHERE id=?", (entrega_id_sol,)).fetchone()
                    con.close()
                    if entrega:
                        st.session_state[f"acuse_ver_{sid}"] = {"folio":entrega_id_sol,"matricula":matricula,"nombre":nombre_emp,
                            "puesto":puesto_emp,"area":area_emp,"disciplina":disciplina_emp,"prenda":nombre_pre,"talla":talla_pre,
                            "cantidad":next((h[1] for h in historial if h[0] == entrega_id_sol), cantidad_sol),"fecha":entrega[0],"responsable":entrega[1],"observaciones":entrega[2],"firma":entrega[3]}
                if st.session_state.get(f"acuse_ver_{sid}"):
                    mostrar_acuse(st.session_state[f"acuse_ver_{sid}"])


# ============================================================
# INVENTARIO
# ============================================================

elif opcion == "👕 Inventario":

    st.title("👕 Control de Inventario")

    pestaña1, pestaña2 = st.tabs([
        "➕ Registrar prenda",
        "📦 Existencias"
    ])

    with pestaña1:

        st.subheader("Registrar prenda")

        with st.form("form_prenda"):

            col1, col2 = st.columns(2)

            with col1:

                catalogo_prendas = [
                    "Selecciona una prenda",
                    "Overall",
                    "Botas MR",
                    "Botas DURAMAX",
                    "Botas PETROLERAS",
                    "Otra prenda (captura manual)"
                ]
                seleccion_prenda = st.selectbox(
                    "Prenda del catálogo *", catalogo_prendas
                )
                if seleccion_prenda == "Otra prenda (captura manual)":
                    nombre = st.text_input("Nombre de la prenda *")
                elif seleccion_prenda == "Selecciona una prenda":
                    nombre = ""
                else:
                    nombre = seleccion_prenda

                tipo = st.selectbox(
                    "Tipo",
                    ["Uniforme", "Accesorio"]
                )

                if seleccion_prenda == "Overall":
                    talla = st.selectbox(
                        "Talla de overall *",
                        ["Selecciona una talla"] + [str(n) for n in range(28, 53, 2)]
                    )
                elif seleccion_prenda in ("Botas MR", "Botas DURAMAX", "Botas PETROLERAS"):
                    talla = st.selectbox(
                        "Número de botas *",
                        ["Selecciona un número"] + [str(n) for n in range(22, 32)]
                    )
                else:
                    talla = st.text_input("Talla *")

            with col2:

                existencia = st.number_input(
                    "Existencia inicial",
                    min_value=0,
                    value=0,
                    step=1
                )

                minimo = st.number_input(
                    "Inventario mínimo",
                    min_value=0,
                    value=0,
                    step=1
                )

            guardar = st.form_submit_button(
                "💾 Guardar prenda",
                use_container_width=True
            )

            if guardar:

                if not nombre.strip() or not talla.strip() or talla.startswith("Selecciona"):

                    st.error(
                        "❌ Nombre y talla son obligatorios."
                    )

                else:

                    conexion = conectar()
                    cursor = conexion.cursor()

                    try:

                        cursor.execute("""
                            INSERT INTO prendas
                            (nombre, tipo, talla, existencia, minimo)
                            VALUES (?, ?, ?, ?, ?)
                        """, (
                            nombre.strip(),
                            tipo,
                            talla.strip(),
                            existencia,
                            minimo
                        ))

                        conexion.commit()
                        registrar_bitacora("Inventario", "Alta de prenda", f"{nombre.strip()} - Talla {talla.strip()}")

                        st.success(
                            "✅ Prenda registrada correctamente."
                        )

                    except sqlite3.IntegrityError:

                        st.error(
                            "❌ Ya existe esa prenda con esa talla."
                        )

                    finally:

                        conexion.close()

    with pestaña2:

        st.subheader("📦 Existencias actuales")

        conexion = conectar()
        cursor = conexion.cursor()

        cursor.execute("""
            SELECT
                id,
                nombre,
                tipo,
                talla,
                existencia,
                minimo
            FROM prendas
            WHERE activo = 1
            ORDER BY nombre, talla
        """)

        prendas = cursor.fetchall()
        conexion.close()

        if prendas:

            import pandas as pd

            df = pd.DataFrame(
                prendas,
                columns=[
                    "ID",
                    "Prenda",
                    "Tipo",
                    "Talla",
                    "Existencia",
                    "Mínimo"
                ]
            )

            st.dataframe(
                df,
                use_container_width=True,
                hide_index=True
            )

            st.divider()

            st.subheader("⚠️ Alertas")

            alertas = [
                p for p in prendas
                if p[4] <= p[5]
            ]

            if alertas:

                for p in alertas:

                    st.warning(
                        f"⚠️ {p[1]} talla {p[3]}: "
                        f"{p[4]} disponibles. "
                        f"Mínimo: {p[5]}."
                    )

            else:

                st.success(
                    "✅ No hay prendas por debajo del mínimo."
                )

        else:

            st.info("No hay prendas registradas.")


# ============================================================
# ALMACÉN: ENTRADAS, SALIDAS Y FIRMA
# ============================================================

elif opcion == "🏭 Almacén":

    st.title("🏭 Almacén")
    st.write("Control de entradas, salidas y firma de trabajadores.")

    tab_ent, tab_sal, tab_mov = st.tabs([
        "📥 Entradas",
        "📤 Salidas / Entregas",
        "📋 Movimientos"
    ])

    with tab_ent:
        st.subheader("📥 Entrada de mercancía")
        conexion = conectar(); cursor = conexion.cursor()
        cursor.execute("SELECT id, nombre, talla, existencia FROM prendas WHERE activo = 1 ORDER BY nombre, talla")
        prendas = cursor.fetchall(); conexion.close()

        if not prendas:
            st.warning("⚠️ Primero registra prendas en Inventario.")
        else:
            opciones_prendas = {f"{p[1]} - Talla {p[2]} - Existencia: {p[3]}": p for p in prendas}
            with st.form("form_entrada_almacen"):
                prenda_txt = st.selectbox("Prenda *", list(opciones_prendas.keys()))
                cantidad_ent = st.number_input("Cantidad recibida *", min_value=1, value=1, step=1)
                fecha_ent = st.date_input("Fecha de entrada")
                referencia_ent = st.text_input("Referencia / folio")
                responsable_ent = st.text_input("Responsable de recepción *", value=st.session_state.usuario)
                obs_ent = st.text_area("Observaciones")
                guardar_ent = st.form_submit_button("📥 Registrar entrada", use_container_width=True)
            if guardar_ent:
                prenda = opciones_prendas[prenda_txt]
                if not responsable_ent.strip():
                    st.error("❌ Indica el responsable de recepción.")
                else:
                    conexion = conectar(); cursor = conexion.cursor()
                    try:
                        fecha_hora = f"{fecha_ent} {datetime.now().strftime('%H:%M:%S')}"
                        nueva = prenda[3] + cantidad_ent
                        cursor.execute("UPDATE prendas SET existencia = ? WHERE id = ?", (nueva, prenda[0]))
                        desc = f"Entrada de almacén | Ref: {referencia_ent.strip() or 'Sin referencia'} | {obs_ent.strip()}"
                        cursor.execute("""INSERT INTO movimientos (prenda_id,tipo,cantidad,fecha,usuario,descripcion) VALUES (?, 'ENTRADA', ?, ?, ?, ?)""", (prenda[0], cantidad_ent, fecha_hora, responsable_ent.strip(), desc))
                        conexion.commit()
                        st.success(f"✅ Entrada registrada. Existencia actual: {nueva}.")
                    except Exception as e:
                        conexion.rollback(); st.error(f"❌ Error: {e}")
                    finally: conexion.close()

    with tab_sal:
        st.subheader("📤 Salida de almacén / entrega al trabajador")
        if st_canvas is None:
            st.error("⚠️ Falta la función de firma digital.")
            st.code("python -m pip install streamlit-drawable-canvas pillow reportlab", language="powershell")
        else:
            conexion = conectar(); cursor = conexion.cursor()
            cursor.execute("SELECT id, numero_empleado, nombre, puesto, area, disciplina, COALESCE(tipo_trabajador, 'Sin clasificar') FROM empleados WHERE activo = 1 ORDER BY nombre")
            empleados = cursor.fetchall()
            cursor.execute("SELECT id, nombre, talla, existencia FROM prendas WHERE activo = 1 ORDER BY nombre, talla")
            prendas = cursor.fetchall(); conexion.close()
            if not empleados:
                st.warning("⚠️ Primero registra trabajadores.")
            elif not prendas:
                st.warning("⚠️ Primero registra prendas.")
            else:
                empleados_d = {f"{e[1]} - {e[2]} ({e[6]})": e for e in empleados}
                prendas_d = {f"{p[1]} - Talla {p[2]} - Disponible: {p[3]}": p for p in prendas}
                empleado_txt = st.selectbox("Trabajador (Matrícula) *", list(empleados_d.keys()), key="alm_emp")
                empleado = empleados_d[empleado_txt]
                prenda_txt = st.selectbox("Prenda *", list(prendas_d.keys()), key="alm_prenda")
                prenda = prendas_d[prenda_txt]
                c1,c2 = st.columns(2)
                with c1: cantidad_sal = st.number_input("Cantidad *", min_value=1, value=1, step=1, key="alm_cant")
                with c2: fecha_sal = st.date_input("Fecha de salida", key="alm_fecha")
                responsable_sal = st.text_input("Responsable de entrega *", value=st.session_state.usuario, key="alm_resp")
                obs_sal = st.text_area("Observaciones", key="alm_obs")
                st.markdown("### ✍️ Firma del trabajador")
                st.info("El trabajador debe firmar antes de confirmar la salida.")
                canvas = st_canvas(fill_color="rgba(255,255,255,0)", stroke_width=3, stroke_color="#111111", background_color="#FFFFFF", height=220, width=700, drawing_mode="freedraw", display_toolbar=True, key="firma_almacen")
                confirmar = st.button("✅ Registrar salida y guardar firma", use_container_width=True, key="alm_confirmar")
                if confirmar:
                    firma = firma_a_base64(canvas.image_data if canvas else None)
                    if not responsable_sal.strip(): st.error("❌ Indica el responsable de entrega.")
                    elif not firma: st.error("❌ El trabajador debe realizar su firma.")
                    elif prenda[3] < cantidad_sal: st.error(f"❌ Inventario insuficiente. Disponible: {prenda[3]}.")
                    else:
                        conexion = conectar(); cursor = conexion.cursor()
                        try:
                            fecha_hora = f"{fecha_sal} {datetime.now().strftime('%H:%M:%S')}"
                            nueva = prenda[3] - cantidad_sal
                            cursor.execute("UPDATE prendas SET existencia = ? WHERE id = ?", (nueva, prenda[0]))
                            cursor.execute("""INSERT INTO entregas (empleado_id,prenda_id,cantidad,fecha,responsable,acuse,observaciones,firma,firma_fecha) VALUES (?,?,?, ?,?,1,?,?,?)""", (empleado[0],prenda[0],cantidad_sal,fecha_hora,responsable_sal.strip(),obs_sal.strip(),firma,fecha_hora))
                            folio = cursor.lastrowid
                            cursor.execute("""INSERT INTO movimientos (prenda_id,tipo,cantidad,fecha,usuario,descripcion) VALUES (?, 'SALIDA', ?, ?, ?, ?)""", (prenda[0],cantidad_sal,fecha_hora,st.session_state.usuario,"Salida de almacén con firma del trabajador"))
                            conexion.commit()
                            st.session_state.ultimo_acuse={"folio":folio,"matricula":empleado[1],"nombre":empleado[2],"puesto":empleado[3],"area":empleado[4],"disciplina":empleado[5] or "","prenda":prenda[1],"talla":prenda[2],"cantidad":cantidad_sal,"fecha":fecha_hora,"responsable":responsable_sal.strip(),"observaciones":obs_sal.strip(),"firma":firma}
                            st.success("✅ Salida registrada y firma guardada.")
                            st.info(f"Existencia actual: {nueva}")
                            st.rerun()
                        except Exception as e:
                            conexion.rollback(); st.error(f"❌ Error: {e}")
                        finally: conexion.close()
                if st.session_state.get("ultimo_acuse") is not None:
                    st.divider(); mostrar_acuse(st.session_state.ultimo_acuse)

    with tab_mov:
        st.subheader("📋 Historial de movimientos")
        conexion=conectar(); cursor=conexion.cursor()
        cursor.execute("""SELECT m.id,m.fecha,m.tipo,p.nombre,p.talla,m.cantidad,m.usuario,m.descripcion FROM movimientos m INNER JOIN prendas p ON m.prenda_id=p.id ORDER BY m.fecha DESC""")
        movs=cursor.fetchall(); conexion.close()
        if movs:
            import pandas as pd
            st.dataframe(pd.DataFrame(movs,columns=["Folio","Fecha","Tipo","Prenda","Talla","Cantidad","Usuario","Descripción"]),use_container_width=True,hide_index=True)
        else: st.info("No hay movimientos registrados.")


# ============================================================
# ENTREGAS
# ============================================================
elif opcion == "📦 Entregas":

    st.title("📦 Registro de Entregas de Uniformes")
    st.write("Selecciona al empleado, registra el uniforme y captura la firma del empleado.")

    if st_canvas is None:
        st.error("⚠️ Falta instalar la función de firma digital.")
        st.code("python -m pip install streamlit-drawable-canvas pillow reportlab", language="powershell")
        st.stop()

    conexion = conectar()
    cursor = conexion.cursor()
    cursor.execute("SELECT id, numero_empleado, nombre, puesto, area, COALESCE(tipo_trabajador, 'Sin clasificar') FROM empleados WHERE activo = 1 ORDER BY nombre")
    empleados = cursor.fetchall()
    cursor.execute("SELECT id, nombre, talla, existencia FROM prendas WHERE activo = 1 ORDER BY nombre, talla")
    prendas = cursor.fetchall()
    conexion.close()

    if not empleados:
        st.warning("⚠️ Primero registra empleados.")
    elif not prendas:
        st.warning("⚠️ Primero registra prendas en inventario.")
    else:
        empleados_dict = {f"{e[1]} - {e[2]} ({e[5]})": e for e in empleados}
        prendas_dict = {f"{p[1]} - Talla {p[2]} - Disponible: {p[3]}": p for p in prendas}

        izquierda, derecha = st.columns([1.05, 0.95])

        with izquierda:
            st.subheader("📋 Datos de la entrega")
            empleado_texto = st.selectbox("Empleado (Matrícula) *", list(empleados_dict.keys()))
            empleado = empleados_dict[empleado_texto]
            prenda_texto = st.selectbox("Prenda *", list(prendas_dict.keys()))
            prenda = prendas_dict[prenda_texto]

            c1, c2 = st.columns(2)
            with c1:
                cantidad = st.number_input("Cantidad *", min_value=1, value=1, step=1)
            with c2:
                fecha_entrega = st.date_input("Fecha de entrega")

            responsable = st.text_input("Responsable de entrega *", value=st.session_state.usuario)
            observaciones = st.text_area("Observaciones")

            st.subheader("✍️ Firma del empleado")
            st.info("El empleado debe firmar dentro del recuadro.")
            canvas_result = st_canvas(
                fill_color="rgba(255, 255, 255, 0)",
                stroke_width=3,
                stroke_color="#111111",
                background_color="#FFFFFF",
                height=220,
                width=700,
                drawing_mode="freedraw",
                display_toolbar=True,
                key="firma_entrega"
            )

            confirmar = st.button("✅ Confirmar entrega y guardar firma", use_container_width=True)

            if confirmar:
                datos_firma = canvas_result.image_data if canvas_result else None
                firma = firma_a_base64(datos_firma)
                if not responsable.strip():
                    st.error("❌ Indica el responsable de entrega.")
                elif not firma:
                    st.error("❌ El empleado debe realizar su firma.")
                elif prenda[3] < cantidad:
                    st.error(f"❌ Inventario insuficiente. Disponible: {prenda[3]}.")
                else:
                    conexion = conectar()
                    cursor = conexion.cursor()
                    try:
                        fecha_hora = f"{fecha_entrega} {datetime.now().strftime('%H:%M:%S')}"
                        nueva_existencia = prenda[3] - cantidad
                        cursor.execute("UPDATE prendas SET existencia = ? WHERE id = ?", (nueva_existencia, prenda[0]))
                        cursor.execute("""
                            INSERT INTO entregas
                            (empleado_id, prenda_id, cantidad, fecha, responsable, acuse, observaciones, firma, firma_fecha)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (empleado[0], prenda[0], cantidad, fecha_hora, responsable.strip(), 1, observaciones.strip(), firma, fecha_hora))
                        cursor.execute("""
                            INSERT INTO movimientos
                            (prenda_id, tipo, cantidad, fecha, usuario, descripcion)
                            VALUES (?, ?, ?, ?, ?, ?)
                        """, (prenda[0], "SALIDA", cantidad, fecha_hora, st.session_state.usuario, "Entrega de uniforme con firma digital"))
                        folio = cursor.lastrowid
                        conexion.commit()

                        datos = {
                            "folio": folio,
                            "matricula": empleado[1],
                            "nombre": empleado[2],
                            "puesto": empleado[3],
                            "area": empleado[4],
                            "prenda": prenda[1],
                            "talla": prenda[2],
                            "cantidad": cantidad,
                            "fecha": fecha_hora,
                            "responsable": responsable.strip(),
                            "observaciones": observaciones.strip(),
                            "firma": firma,
                        }
                        st.session_state.ultimo_acuse = datos
                        st.session_state.firma_key = st.session_state.get("firma_key", 0) + 1
                        st.success("✅ Entrega registrada correctamente.")
                        st.rerun()
                    except Exception as error:
                        conexion.rollback()
                        st.error(f"❌ Error al registrar la entrega: {error}")
                    finally:
                        conexion.close()

        with derecha:
            st.subheader("🧾 Acuse de entrega")
            if st.session_state.get("ultimo_acuse") is not None:
                mostrar_acuse(st.session_state.ultimo_acuse)
            else:
                st.info("Después de confirmar la entrega y la firma, aquí aparecerá el acuse digital. No es necesario imprimirlo.")

        st.divider()
        st.subheader("📋 Historial de entregas")
        conexion = conectar(); cursor = conexion.cursor()
        cursor.execute("SELECT DISTINCT talla FROM prendas WHERE activo=1 ORDER BY talla")
        tallas = [x[0] for x in cursor.fetchall() if x[0]]
        cursor.execute("SELECT DISTINCT disciplina FROM empleados WHERE activo=1 AND disciplina IS NOT NULL AND disciplina != '' ORDER BY disciplina")
        disciplinas = [x[0] for x in cursor.fetchall()]
        conexion.close()
        f1,f2,f3=st.columns(3)
        with f1: fecha_f = st.date_input("📅 Fecha", value=None, key="f_fecha")
        with f2: talla_f = st.selectbox("👕 Talla", ["Todas"]+tallas, key="f_talla")
        with f3: disciplina_f = st.selectbox("🎯 Disciplina", ["Todas"]+disciplinas, key="f_disciplina")
        conexion=conectar(); cursor=conexion.cursor()
        q="""SELECT en.id,e.numero_empleado,e.nombre,COALESCE(e.disciplina,''),p.nombre,p.talla,en.cantidad,en.fecha,en.responsable,CASE WHEN en.firma IS NOT NULL AND en.firma!='' THEN 'Sí' ELSE 'No' END FROM entregas en INNER JOIN empleados e ON en.empleado_id=e.id INNER JOIN prendas p ON en.prenda_id=p.id WHERE 1=1"""
        params=[]
        if fecha_f: q += " AND date(en.fecha)=?"; params.append(str(fecha_f))
        if talla_f!="Todas": q += " AND p.talla=?"; params.append(talla_f)
        if disciplina_f!="Todas": q += " AND e.disciplina=?"; params.append(disciplina_f)
        q += " ORDER BY en.fecha DESC"
        cursor.execute(q,params); historial=cursor.fetchall(); conexion.close()
        if historial:
            import pandas as pd
            df=pd.DataFrame(historial,columns=["Folio","Matrícula","Empleado","Disciplina","Prenda","Talla","Cantidad","Fecha","Responsable","Firma"])
            st.metric("📊 Registros encontrados",len(df))
            st.dataframe(df,use_container_width=True,hide_index=True)
        else: st.info("No hay entregas que coincidan con los filtros.")


# ============================================================
# ACUSES Y FIRMAS
# ============================================================

elif opcion == "✍️ Acuses y firmas":

    st.title("✍️ Acuses y firmas")
    st.write("Consulta las entregas y descarga el acuse con la firma del empleado.")

    conexion = conectar()
    cursor = conexion.cursor()
    cursor.execute("""
        SELECT en.id, e.numero_empleado, e.nombre, e.puesto, e.area, e.disciplina,
               p.nombre, p.talla, en.cantidad, en.fecha, en.responsable,
               en.observaciones, en.firma
        FROM entregas en
        INNER JOIN empleados e ON en.empleado_id = e.id
        INNER JOIN prendas p ON en.prenda_id = p.id
        ORDER BY en.fecha DESC
    """)
    filas = cursor.fetchall()
    conexion.close()

    if not filas:
        st.info("No hay entregas registradas.")
    else:
        opciones_acuse = {f"ENT-{x[0]:06d} | {x[1]} - {x[2]}": x for x in filas}
        seleccion = st.selectbox("Selecciona una entrega", list(opciones_acuse.keys()))
        x = opciones_acuse[seleccion]
        datos = {
            "folio": x[0], "matricula": x[1], "nombre": x[2], "puesto": x[3],
            "area": x[4], "disciplina": x[5], "prenda": x[6], "talla": x[7], "cantidad": x[8],
            "fecha": x[9], "responsable": x[10], "observaciones": x[11], "firma": x[12]
        }
        mostrar_acuse(datos)


# ============================================================
# REPOSICIONES
# ============================================================

elif opcion == "🔄 Reposiciones":

    st.title("🔄 Gestión de Reposiciones")

    conexion = conectar()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT id, numero_empleado, nombre
        FROM empleados
        WHERE activo = 1
        ORDER BY nombre
    """)

    empleados = cursor.fetchall()

    cursor.execute("""
        SELECT id, nombre, talla, existencia
        FROM prendas
        WHERE activo = 1
        ORDER BY nombre, talla
    """)

    prendas = cursor.fetchall()

    conexion.close()

    if not empleados:

        st.warning("⚠️ Registra empleados primero.")

    elif not prendas:

        st.warning("⚠️ Registra prendas primero.")

    else:

        st.subheader("📝 Nueva solicitud")

        with st.form("form_reposicion"):

            empleados_dict = {
                f"{e[1]} - {e[2]}": e[0]
                for e in empleados
            }

            empleado_texto = st.selectbox(
                "Empleado",
                list(empleados_dict.keys())
            )

            prendas_dict = {
                f"{p[1]} - Talla {p[2]} - "
                f"Disponible: {p[3]}": p[0]
                for p in prendas
            }

            prenda_texto = st.selectbox(
                "Prenda",
                list(prendas_dict.keys())
            )

            cantidad = st.number_input(
                "Cantidad",
                min_value=1,
                value=1,
                step=1
            )

            with conectar() as cm:
                motivos_disponibles = [r[0] for r in cm.execute(
                    "SELECT motivo FROM catalogo_motivos_reposicion WHERE activo=1 ORDER BY motivo")]
            motivo = st.selectbox("Motivo", motivos_disponibles)

            observaciones = st.text_area(
                "Observaciones"
            )

            solicitar = st.form_submit_button(
                "🔄 Solicitar reposición",
                use_container_width=True
            )

            if solicitar:

                empleado_id = empleados_dict[empleado_texto]
                prenda_id = prendas_dict[prenda_texto]

                conexion = conectar()
                cursor = conexion.cursor()

                cursor.execute("""
                    INSERT INTO reposiciones
                    (empleado_id, prenda_id, cantidad, motivo,
                     fecha_solicitud, estado, observaciones)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (
                    empleado_id,
                    prenda_id,
                    cantidad,
                    motivo,
                    ahora(),
                    "Pendiente",
                    observaciones.strip()
                ))

                conexion.commit()
                conexion.close()
                registrar_bitacora("Reposiciones", "Solicitud de reposición", f"Empleado ID: {empleado_id} | Prenda ID: {prenda_id} | Cantidad: {cantidad}")

                st.success(
                    "✅ Solicitud registrada como Pendiente."
                )

        st.divider()

        st.subheader("🟡 Solicitudes pendientes")

        conexion = conectar()
        cursor = conexion.cursor()

        cursor.execute("""
            SELECT
                r.id,
                r.empleado_id,
                r.prenda_id,
                e.numero_empleado,
                e.nombre,
                p.nombre,
                p.talla,
                r.cantidad,
                r.motivo,
                r.fecha_solicitud,
                r.observaciones
            FROM reposiciones r
            INNER JOIN empleados e
                ON r.empleado_id = e.id
            INNER JOIN prendas p
                ON r.prenda_id = p.id
            WHERE r.estado = 'Pendiente'
            ORDER BY r.fecha_solicitud DESC
        """)

        pendientes = cursor.fetchall()
        conexion.close()

        if not pendientes:

            st.info(
                "No existen solicitudes pendientes."
            )

        else:

            for solicitud in pendientes:

                (
                    reposicion_id,
                    empleado_id,
                    prenda_id,
                    numero_empleado,
                    nombre_empleado,
                    nombre_prenda,
                    talla,
                    cantidad,
                    motivo,
                    fecha_solicitud,
                    observaciones
                ) = solicitud

                with st.expander(
                    f"Solicitud #{reposicion_id} - "
                    f"{numero_empleado} - {nombre_empleado}"
                ):

                    st.write(
                        f"**Prenda:** {nombre_prenda}"
                    )

                    st.write(
                        f"**Talla:** {talla}"
                    )

                    st.write(
                        f"**Cantidad:** {cantidad}"
                    )

                    st.write(
                        f"**Motivo:** {motivo}"
                    )

                    st.write(
                        f"**Fecha:** {fecha_solicitud}"
                    )

                    if observaciones:

                        st.write(
                            f"**Observaciones:** {observaciones}"
                        )

                    col1, col2 = st.columns(2)

                    with col1:

                        autorizar = st.button(
                            "✅ Autorizar",
                            key=f"autorizar_{reposicion_id}"
                        )

                    with col2:

                        rechazar = st.button(
                            "❌ Rechazar",
                            key=f"rechazar_{reposicion_id}"
                        )

                    if autorizar:

                        conexion = conectar()
                        cursor = conexion.cursor()

                        try:

                            cursor.execute("""
                                SELECT existencia, minimo
                                FROM prendas
                                WHERE id = ?
                            """, (prenda_id,))

                            inventario = cursor.fetchone()

                            if inventario is None:

                                st.error(
                                    "❌ No se encontró la prenda."
                                )

                            elif inventario[0] < cantidad:

                                st.error(
                                    f"❌ No hay inventario suficiente. "
                                    f"Disponible: {inventario[0]}. "
                                    f"Solicitado: {cantidad}."
                                )

                            else:

                                nueva_existencia = (
                                    inventario[0] - cantidad
                                )

                                cursor.execute("""
                                    UPDATE prendas
                                    SET existencia = ?
                                    WHERE id = ?
                                """, (
                                    nueva_existencia,
                                    prenda_id
                                ))

                                cursor.execute("""
                                    UPDATE reposiciones
                                    SET
                                        fecha_entrega = ?,
                                        autorizado_por = ?,
                                        entregado_por = ?,
                                        estado = ?
                                    WHERE id = ?
                                """, (
                                    ahora(),
                                    st.session_state.usuario,
                                    st.session_state.usuario,
                                    "Autorizada",
                                    reposicion_id
                                ))

                                cursor.execute("""
                                    INSERT INTO movimientos
                                    (prenda_id, tipo, cantidad, fecha,
                                     usuario, descripcion)
                                    VALUES (?, ?, ?, ?, ?, ?)
                                """, (
                                    prenda_id,
                                    "REPOSICION",
                                    cantidad,
                                    ahora(),
                                    st.session_state.usuario,
                                    f"Reposición autorizada. Motivo: {motivo}"
                                ))

                                conexion.commit()
                                registrar_bitacora("Reposiciones", "Reposición autorizada", f"Solicitud: {reposicion_id} | Prenda ID: {prenda_id} | Cantidad: {cantidad}")

                                st.success(
                                    "✅ Reposición autorizada y "
                                    "descontada del inventario."
                                )

                                if nueva_existencia <= inventario[1]:

                                    st.warning(
                                        "⚠️ El inventario llegó "
                                        "al nivel mínimo."
                                    )

                        except Exception as error:

                            conexion.rollback()

                            st.error(
                                f"❌ Error: {error}"
                            )

                        finally:

                            conexion.close()

                        st.rerun()

                    if rechazar:

                        conexion = conectar()
                        cursor = conexion.cursor()

                        cursor.execute("""
                            UPDATE reposiciones
                            SET
                                estado = ?,
                                autorizado_por = ?
                            WHERE id = ?
                        """, (
                            "Rechazada",
                            st.session_state.usuario,
                            reposicion_id
                        ))

                        conexion.commit()
                        conexion.close()
                        registrar_bitacora("Reposiciones", "Reposición rechazada", f"Solicitud: {reposicion_id}")

                        st.success(
                            "❌ Solicitud rechazada."
                        )

                        st.rerun()

        st.divider()

        st.subheader("📋 Historial de reposiciones")

        conexion = conectar()
        cursor = conexion.cursor()

        cursor.execute("""
            SELECT
                e.numero_empleado,
                e.nombre,
                p.nombre,
                p.talla,
                r.cantidad,
                r.motivo,
                r.fecha_solicitud,
                r.fecha_entrega,
                r.autorizado_por,
                r.entregado_por,
                r.estado
            FROM reposiciones r
            INNER JOIN empleados e
                ON r.empleado_id = e.id
            INNER JOIN prendas p
                ON r.prenda_id = p.id
            ORDER BY r.fecha_solicitud DESC
        """)

        historial = cursor.fetchall()
        conexion.close()

        if historial:

            import pandas as pd

            df = pd.DataFrame(
                historial,
                columns=[
                    "Matrícula",
                    "Empleado",
                    "Prenda",
                    "Talla",
                    "Cantidad",
                    "Motivo",
                    "Solicitud",
                    "Entrega",
                    "Autorizó",
                    "Entregó",
                    "Estado"
                ]
            )

            st.dataframe(
                df,
                use_container_width=True,
                hide_index=True
            )

        else:

            st.info(
                "No existen reposiciones registradas."
            )


# ============================================================
# MATRIZ DE DOTACIÓN
# ============================================================

elif opcion == "📋 Matriz de Dotación":

    st.title("📋 Matriz de Dotación")

    st.write(
        "Permite definir qué prenda corresponde a cada puesto, "
        "cantidad y frecuencia."
    )

    conexion = conectar()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT id, nombre, talla
        FROM prendas
        WHERE activo = 1
        ORDER BY nombre, talla
    """)

    prendas = cursor.fetchall()
    conexion.close()

    if not prendas:

        st.warning(
            "⚠️ Registra prendas en Inventario antes "
            "de crear la matriz."
        )

    else:

        with st.form("form_dotacion"):

            puesto = st.text_input(
                "Puesto *"
            )

            tipo_dotacion = st.selectbox(
                "Tipo de trabajador al que aplica *",
                ["Todos", "Sindicalizado", "Base"],
                help="Usa 'Todos' solo si la misma regla aplica a ambos grupos."
            )

            prendas_dict = {
                f"{p[1]} - Talla {p[2]}": p[0]
                for p in prendas
            }

            prenda_texto = st.selectbox(
                "Prenda",
                list(prendas_dict.keys())
            )

            cantidad = st.number_input(
                "Cantidad autorizada",
                min_value=1,
                value=1,
                step=1
            )

            frecuencia = st.selectbox(
                "Frecuencia",
                [
                    "Al ingreso",
                    "Anual",
                    "Semestral",
                    "Según política",
                    "Otro"
                ]
            )

            observaciones = st.text_area(
                "Observaciones"
            )

            guardar = st.form_submit_button(
                "💾 Guardar regla de dotación",
                use_container_width=True
            )

            if guardar:

                if not puesto.strip():

                    st.error(
                        "❌ Indica el puesto."
                    )

                else:

                    conexion = conectar()
                    cursor = conexion.cursor()

                    cursor.execute("""
                        INSERT INTO dotaciones
                        (puesto, prenda_id, cantidad,
                         frecuencia, observaciones, tipo_trabajador)
                        VALUES (?, ?, ?, ?, ?, ?)
                    """, (
                        puesto.strip(),
                        prendas_dict[prenda_texto],
                        cantidad,
                        frecuencia,
                        observaciones.strip(),
                        tipo_dotacion
                    ))

                    conexion.commit()
                    conexion.close()
                    registrar_bitacora("Matriz de Dotación", "Alta de regla", f"Puesto: {puesto.strip()} | Tipo: {tipo_dotacion} | Prenda ID: {prendas_dict[prenda_texto]} | Cantidad: {cantidad} | Frecuencia: {frecuencia}")

                    st.success(
                        "✅ Regla de dotación guardada."
                    )

        st.divider()

        st.subheader("📋 Reglas registradas")

        conexion = conectar()
        cursor = conexion.cursor()

        cursor.execute("""
            SELECT
                d.puesto,
                p.nombre,
                p.talla,
                d.cantidad,
                d.frecuencia,
                d.observaciones,
                COALESCE(d.tipo_trabajador, 'Todos')
            FROM dotaciones d
            INNER JOIN prendas p
                ON d.prenda_id = p.id
            ORDER BY d.puesto, p.nombre
        """)

        datos = cursor.fetchall()
        conexion.close()

        if datos:

            import pandas as pd

            df = pd.DataFrame(
                datos,
                columns=[
                    "Puesto",
                    "Prenda",
                    "Talla",
                    "Cantidad",
                    "Frecuencia",
                    "Observaciones",
                    "Tipo de trabajador"
                ]
            )

            st.dataframe(
                df,
                use_container_width=True,
                hide_index=True
            )

        else:

            st.info(
                "No hay reglas de dotación registradas."
            )


# ============================================================
# FRECUENCIAS
# ============================================================

elif opcion == "📈 Frecuencias":
    st.title("📈 Frecuencias")
    st.write("Analiza cuántas entregas se realizan por mes, disciplina y talla, y consulta las frecuencias definidas en la matriz de dotación.")
    import pandas as pd
    conexion=conectar(); cursor=conexion.cursor()
    cursor.execute("SELECT substr(fecha,1,7),COUNT(*),COALESCE(SUM(cantidad),0) FROM entregas GROUP BY substr(fecha,1,7) ORDER BY substr(fecha,1,7) DESC")
    por_mes=cursor.fetchall()
    cursor.execute("SELECT COALESCE(e.disciplina,'Sin disciplina'),COUNT(*),COALESCE(SUM(en.cantidad),0) FROM entregas en INNER JOIN empleados e ON en.empleado_id=e.id GROUP BY COALESCE(e.disciplina,'Sin disciplina') ORDER BY COUNT(*) DESC")
    por_disc=cursor.fetchall()
    cursor.execute("SELECT p.talla,COUNT(*),COALESCE(SUM(en.cantidad),0) FROM entregas en INNER JOIN prendas p ON en.prenda_id=p.id GROUP BY p.talla ORDER BY COUNT(*) DESC")
    por_talla=cursor.fetchall()
    cursor.execute("SELECT frecuencia,COUNT(*),COALESCE(SUM(cantidad),0) FROM dotaciones GROUP BY frecuencia ORDER BY COUNT(*) DESC")
    por_frec=cursor.fetchall()
    conexion.close()
    a,b,c=st.columns(3)
    with a: st.metric("📦 Entregas",sum(x[1] for x in por_mes))
    with b: st.metric("👕 Unidades",sum(x[2] for x in por_mes))
    with c: st.metric("🎯 Disciplinas",len(por_disc))
    if por_mes:
        st.subheader("📅 Frecuencia por mes")
        st.dataframe(pd.DataFrame(por_mes,columns=["Mes","Entregas","Unidades"]),use_container_width=True,hide_index=True)
    c1,c2=st.columns(2)
    with c1:
        st.subheader("🎯 Frecuencia por disciplina")
        if por_disc: st.dataframe(pd.DataFrame(por_disc,columns=["Disciplina","Entregas","Unidades"]),use_container_width=True,hide_index=True)
        else: st.info("No hay disciplinas registradas.")
    with c2:
        st.subheader("👕 Frecuencia por talla")
        if por_talla: st.dataframe(pd.DataFrame(por_talla,columns=["Talla","Entregas","Unidades"]),use_container_width=True,hide_index=True)
        else: st.info("No hay entregas registradas.")
    st.subheader("🔄 Frecuencias de la Matriz de Dotación")
    if por_frec: st.dataframe(pd.DataFrame(por_frec,columns=["Frecuencia","Reglas","Unidades autorizadas"]),use_container_width=True,hide_index=True)
    else: st.info("No hay frecuencias configuradas en la Matriz de Dotación.")


# ============================================================
# DASHBOARD
# ============================================================
elif opcion == "📊 Dashboard":

    st.title("📊 Dashboard TNG")
    st.write(
        "Indicadores de control de dotaciones e inventario."
    )

    conexion = conectar()
    cursor = conexion.cursor()

    # Empleados
    cursor.execute("""
        SELECT COUNT(*)
        FROM empleados
        WHERE activo = 1
    """)
    total_empleados = cursor.fetchone()[0]

    # Prendas
    cursor.execute("""
        SELECT COUNT(*)
        FROM prendas
        WHERE activo = 1
    """)
    total_prendas = cursor.fetchone()[0]

    # Existencias
    cursor.execute("""
        SELECT COALESCE(SUM(existencia), 0)
        FROM prendas
        WHERE activo = 1
    """)
    total_existencias = cursor.fetchone()[0]

    # Entregas
    cursor.execute("""
        SELECT COUNT(*)
        FROM entregas
    """)
    total_entregas = cursor.fetchone()[0]

    # Unidades entregadas
    cursor.execute("""
        SELECT COALESCE(SUM(cantidad), 0)
        FROM entregas
    """)
    unidades_entregadas = cursor.fetchone()[0]

    # Reposiciones
    cursor.execute("""
        SELECT COUNT(*)
        FROM reposiciones
    """)
    total_reposiciones = cursor.fetchone()[0]

    # Pendientes
    cursor.execute("""
        SELECT COUNT(*)
        FROM reposiciones
        WHERE estado = 'Pendiente'
    """)
    pendientes = cursor.fetchone()[0]

    # Rechazadas
    cursor.execute("""
        SELECT COUNT(*)
        FROM reposiciones
        WHERE estado = 'Rechazada'
    """)
    rechazadas = cursor.fetchone()[0]

    # Inventario bajo
    cursor.execute("""
        SELECT COUNT(*)
        FROM prendas
        WHERE activo = 1
        AND existencia <= minimo
    """)
    inventario_bajo = cursor.fetchone()[0]

    # Indicadores de ahorro de papel: se considera 1 hoja evitada por cada entrega digital.
    cursor.execute("SELECT COUNT(*) FROM entregas WHERE firma IS NOT NULL AND firma != ''")
    entregas_digitales = cursor.fetchone()[0]
    hojas_evitadas = entregas_digitales
    porcentaje_digital = (entregas_digitales / total_entregas * 100) if total_entregas else 0

    conexion.close()

    st.subheader("📌 Indicadores principales")

    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.metric(
            "👥 Empleados",
            total_empleados
        )

    with col2:
        st.metric(
            "👕 Prendas",
            total_prendas
        )

    with col3:
        st.metric(
            "📦 Existencias",
            total_existencias
        )

    with col4:
        st.metric(
            "🚚 Entregas",
            total_entregas
        )

    col5, col6, col7, col8 = st.columns(4)

    with col5:
        st.metric(
            "📦 Unidades entregadas",
            unidades_entregadas
        )

    with col6:
        st.metric(
            "🔄 Reposiciones",
            total_reposiciones
        )

    with col7:
        st.metric(
            "🟡 Pendientes",
            pendientes
        )

    with col8:
        st.metric(
            "⚠️ Inventario bajo",
            inventario_bajo
        )

    st.divider()

    st.subheader("♻️ Control sin papel")
    p1, p2, p3 = st.columns(3)
    with p1:
        st.metric("📄 Hojas evitadas", hojas_evitadas)
    with p2:
        st.metric("✍️ Acuses digitales", entregas_digitales)
    with p3:
        st.metric("♻️ Entregas digitales", f"{porcentaje_digital:.0f}%")
    st.caption("Cálculo estimado: 1 hoja evitada por cada entrega que cuenta con firma digital. El PDF se descarga únicamente cuando sea necesario.")

    st.divider()

    # Alertas
    st.subheader("⚠️ Alertas de inventario")

    conexion = conectar()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            nombre,
            talla,
            existencia,
            minimo
        FROM prendas
        WHERE activo = 1
        AND existencia <= minimo
        ORDER BY existencia ASC
    """)

    alertas = cursor.fetchall()
    conexion.close()

    if alertas:

        for alerta in alertas:

            st.warning(
                f"⚠️ **{alerta[0]} - Talla {alerta[1]}** | "
                f"Existencia: {alerta[2]} | "
                f"Mínimo: {alerta[3]}"
            )

    else:

        st.success(
            "✅ No existen alertas de inventario."
        )

    st.divider()

    # Inventario
    st.subheader("👕 Existencias por prenda")

    conexion = conectar()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            nombre,
            talla,
            existencia,
            minimo
        FROM prendas
        WHERE activo = 1
        ORDER BY nombre, talla
    """)

    inventario = cursor.fetchall()
    conexion.close()

    if inventario:

        import pandas as pd

        df_inventario = pd.DataFrame(
            inventario,
            columns=[
                "Prenda",
                "Talla",
                "Existencia",
                "Mínimo"
            ]
        )

        st.dataframe(
            df_inventario,
            use_container_width=True,
            hide_index=True
        )

    else:

        st.info(
            "No hay datos de inventario."
        )

    st.divider()

    # Gráfica
    st.subheader("📊 Existencias por prenda")

    conexion = conectar()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            nombre,
            SUM(existencia)
        FROM prendas
        WHERE activo = 1
        GROUP BY nombre
        ORDER BY SUM(existencia) DESC
    """)

    grafica = cursor.fetchall()
    conexion.close()

    if grafica:

        import pandas as pd

        df_grafica = pd.DataFrame(
            grafica,
            columns=[
                "Prenda",
                "Existencia"
            ]
        ).set_index("Prenda")

        st.bar_chart(df_grafica)

    else:

        st.info(
            "No hay datos para mostrar la gráfica."
        )

    st.divider()

    # Últimas entregas
    st.subheader("📦 Últimas entregas")

    conexion = conectar()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            e.numero_empleado,
            e.nombre,
            p.nombre,
            p.talla,
            en.cantidad,
            en.fecha,
            en.responsable
        FROM entregas en
        INNER JOIN empleados e
            ON en.empleado_id = e.id
        INNER JOIN prendas p
            ON en.prenda_id = p.id
        ORDER BY en.fecha DESC
        LIMIT 10
    """)

    ultimas = cursor.fetchall()
    conexion.close()

    if ultimas:

        import pandas as pd

        df_ultimas = pd.DataFrame(
            ultimas,
            columns=[
                "Matrícula",
                "Empleado",
                "Prenda",
                "Talla",
                "Cantidad",
                "Fecha",
                "Responsable"
            ]
        )

        st.dataframe(
            df_ultimas,
            use_container_width=True,
            hide_index=True
        )

    else:

        st.info(
            "Todavía no existen entregas."
        )

    st.divider()

    # KPI simple de entregas con acuse
    conexion = conectar()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT COUNT(*)
        FROM entregas
    """)
    total = cursor.fetchone()[0]

    cursor.execute("""
        SELECT COUNT(*)
        FROM entregas
        WHERE acuse = 1
    """)
    con_acuse = cursor.fetchone()[0]

    conexion.close()

    if total > 0:

        porcentaje_acuse = (
            con_acuse / total
        ) * 100

        st.metric(
            "📝 Entregas con acuse",
            f"{porcentaje_acuse:.1f}%"
        )

    else:

        st.metric(
            "📝 Entregas con acuse",
            "0%"
        )


# ============================================================
# KARDEX
# ============================================================
elif opcion == "📦 Kardex":

    st.title("📦 Kardex de movimientos")
    st.write("Consulta el historial de entradas, salidas y reposiciones de cada prenda y talla.")

    conexion = conectar()
    cursor = conexion.cursor()
    cursor.execute("""
        SELECT id, nombre, tipo, talla, existencia, minimo
        FROM prendas
        WHERE activo = 1
        ORDER BY nombre, talla
    """)
    prendas_k = cursor.fetchall()
    conexion.close()

    if not prendas_k:
        st.info("No hay prendas registradas.")
    else:
        opciones_k = {f"{p[1]} - Talla {p[3]}": p for p in prendas_k}
        seleccion_k = st.selectbox("Selecciona una prenda", list(opciones_k.keys()), key="kardex_prenda")
        p = opciones_k[seleccion_k]

        conexion = conectar()
        cursor = conexion.cursor()
        cursor.execute("""
            SELECT id, fecha, tipo, cantidad, usuario, descripcion
            FROM movimientos
            WHERE prenda_id = ?
            ORDER BY datetime(fecha) DESC, id DESC
        """, (p[0],))
        movimientos_k = cursor.fetchall()
        conexion.close()

        st.subheader(f"👕 {p[1]} — Talla {p[3]}")
        c1, c2, c3 = st.columns(3)
        c1.metric("Existencia actual", p[4])
        c2.metric("Stock mínimo", p[5])
        c3.metric("Movimientos", len(movimientos_k))

        if movimientos_k:
            # Reconstrucción aproximada del saldo histórico a partir de la existencia actual.
            saldo = p[4]
            filas = []
            for mov in movimientos_k:
                mov_id, fecha, tipo, cantidad, usuario, descripcion = mov
                saldo_despues = saldo
                if tipo == "ENTRADA":
                    saldo_antes = saldo - cantidad
                elif tipo in ("SALIDA", "REPOSICION"):
                    saldo_antes = saldo + cantidad
                else:
                    saldo_antes = saldo
                filas.append([mov_id, fecha, tipo, cantidad, saldo_antes, saldo_despues, usuario, descripcion or ""])
                saldo = saldo_antes

            import pandas as pd
            df_k = pd.DataFrame(filas, columns=["Folio", "Fecha", "Movimiento", "Cantidad", "Existencia antes", "Existencia después", "Usuario", "Descripción"])
            st.dataframe(df_k, use_container_width=True, hide_index=True)

            st.download_button(
                "📥 Descargar Kardex CSV",
                df_k.to_csv(index=False).encode("utf-8-sig"),
                file_name=f"kardex_{p[1]}_{p[3]}.csv".replace(" ", "_"),
                mime="text/csv",
                use_container_width=True
            )
        else:
            st.info("Esta prenda todavía no tiene movimientos registrados.")


# ============================================================
# ALERTAS
# ============================================================
elif opcion == "⚠️ Alertas":

    st.title("⚠️ Alertas de inventario")
    st.write("Identifica existencias iguales o inferiores al stock mínimo configurado.")

    conexion = conectar()
    cursor = conexion.cursor()
    cursor.execute("""
        SELECT id, nombre, tipo, talla, existencia, minimo
        FROM prendas
        WHERE activo = 1
        ORDER BY existencia ASC, nombre, talla
    """)
    alertas_data = cursor.fetchall()
    conexion.close()

    if not alertas_data:
        st.info("No hay prendas registradas.")
    else:
        criticas = [x for x in alertas_data if x[4] <= 0]
        bajas = [x for x in alertas_data if x[4] > 0 and x[4] <= x[5]]
        normales = [x for x in alertas_data if x[4] > x[5]]

        a1, a2, a3 = st.columns(3)
        a1.metric("🔴 Agotadas", len(criticas))
        a2.metric("🟠 Bajo mínimo", len(bajas))
        a3.metric("🟢 En nivel normal", len(normales))

        if criticas:
            st.subheader("🔴 Agotadas")
            import pandas as pd
            st.dataframe(pd.DataFrame(criticas, columns=["ID","Prenda","Tipo","Talla","Existencia","Mínimo"]), use_container_width=True, hide_index=True)

        if bajas:
            st.subheader("🟠 Existencia baja")
            import pandas as pd
            st.dataframe(pd.DataFrame(bajas, columns=["ID","Prenda","Tipo","Talla","Existencia","Mínimo"]), use_container_width=True, hide_index=True)

        if not criticas and not bajas:
            st.success("✅ Todas las prendas se encuentran por encima del stock mínimo configurado.")

        st.caption("Las alertas se calculan con el campo 'Mínimo' de cada prenda; no modifican el inventario.")


# ============================================================
# NUEVAS FUNCIONES DE CONTROL
# ============================================================
elif opcion == "🚦 Semáforo":
    st.title("🚦 Semáforo de inventario")
    con=conectar()
    filas=con.execute("SELECT nombre,talla,existencia,minimo FROM prendas WHERE activo=1 ORDER BY nombre,talla").fetchall()
    con.close()
    import pandas as pd
    df=pd.DataFrame([{'Prenda':a,'Talla':b,'Existencia':c,'Mínimo':d,'Estado':semaforo(c,d)} for a,b,c,d in filas])
    if df.empty: st.info('Registra prendas para consultar el semáforo.')
    else:
        cols=st.columns(3)
        for col,estado in zip(cols,['🔴 Agotado','🟡 Bajo mínimo','🟢 Disponible']): col.metric(estado,int((df['Estado']==estado).sum()))
        filtro=st.selectbox('Filtrar por estado',['Todos','🔴 Agotado','🟡 Bajo mínimo','🟢 Disponible'])
        st.dataframe(df if filtro=='Todos' else df[df['Estado']==filtro],hide_index=True,use_container_width=True)
    st.caption('Los límites provienen del campo Mínimo configurado en Inventario.')

elif opcion == "📅 Próximas dotaciones":
    st.title('📅 Próxima fecha de dotación')
    st.caption('Cálculo orientativo: anual = 12 meses, semestral = 6 meses; al ingreso se revisa hasta la primera entrega. Las políticas especiales requieren revisión manual.')
    con=conectar(); datos=calcular_proximas_dotaciones(con); con.close()
    import pandas as pd
    df=pd.DataFrame(datos)
    if df.empty: st.info('Registra empleados y reglas de Matriz de Dotación para calcular fechas.')
    else:
        matricula=st.text_input('Buscar por matrícula o empleado')
        if matricula: df=df[df['Matrícula'].astype(str).str.contains(matricula,case=False,regex=False)|df['Empleado'].str.contains(matricula,case=False,regex=False)]
        st.dataframe(df,hide_index=True,use_container_width=True)
        st.metric('Dotaciones vencidas o pendientes',int((df['Estado']=='Vencida / pendiente').sum()))

elif opcion == "📉 Pronóstico":
    st.title('📉 Pronóstico de abastecimiento')
    horizonte=st.slider('Horizonte de planeación (días)',30,180,90,30)
    con=conectar(); proximas=calcular_proximas_dotaciones(con)
    inventario=con.execute('SELECT id,nombre,talla,existencia,minimo FROM prendas WHERE activo=1').fetchall()
    pendientes=con.execute("""SELECT prenda_id,COALESCE(SUM(cantidad-cantidad_entregada),0)
        FROM solicitudes_almacen WHERE estado IN ('Pendiente','Preparada','Parcial') GROUP BY prenda_id""").fetchall()
    con.close()
    pend=dict(pendientes); limite=date.today()+timedelta(days=horizonte)
    # Se agrupa por nombre y talla; se cuenta la demanda programada de matriz, no una predicción estadística.
    demanda={}
    for fila in proximas:
        try: vence=date.fromisoformat(fila['Próxima fecha'])
        except ValueError: continue
        if vence<=limite:
            clave=(fila['Prenda'],fila['Talla']); demanda[clave]=demanda.get(clave,0)+fila['Cantidad']
    import pandas as pd
    filas=[]
    for pid,nombre,talla,existencia,minimo in inventario:
        programada=demanda.get((nombre,talla),0); solicitudes=pend.get(pid,0)
        # Una solicitud puede representar la dotación programada: usar el máximo evita sumarlas dos veces.
        necesidad=max(programada,solicitudes)
        filas.append({'Prenda':nombre,'Talla':talla,'Existencia':existencia,'Mínimo':minimo,
            'Demanda matriz':programada,'Solicitudes pendientes':solicitudes,'Necesidad estimada':necesidad,
            'Compra sugerida':max(0,necesidad+minimo-existencia)})
    df=pd.DataFrame(filas)
    if df.empty: st.info('No hay prendas activas.')
    else:
        st.metric('Unidades sugeridas para abastecer',int(df['Compra sugerida'].sum()))
        st.dataframe(df.sort_values('Compra sugerida',ascending=False),hide_index=True,use_container_width=True)
    st.caption('Estimación de planeación, no orden de compra. Se utiliza el mayor valor entre matriz y solicitudes para evitar una posible doble contabilización; validar manualmente necesidades no coincidentes.')

elif opcion == "↩️ Devoluciones":
    st.title('↩️ Devoluciones de uniformes')
    con=conectar()
    empleados=con.execute('SELECT id,numero_empleado,nombre FROM empleados WHERE activo=1 ORDER BY nombre').fetchall()
    prendas=con.execute('SELECT id,nombre,talla FROM prendas WHERE activo=1 ORDER BY nombre,talla').fetchall()
    con.close()
    if not empleados or not prendas: st.info('Se requieren empleados y prendas registradas.')
    else:
        with st.form('registrar_devolucion'):
            empleado=st.selectbox('Trabajador',empleados,format_func=lambda x:f'{x[1]} - {x[2]}')
            prenda=st.selectbox('Prenda devuelta',prendas,format_func=lambda x:f'{x[1]} - {x[2]}')
            cantidad=st.number_input('Cantidad',min_value=1,step=1)
            motivo=st.selectbox('Motivo',['Baja','Cambio de talla','Reposición','Otra causa'])
            condicion=st.selectbox('Condición',['Reutilizable','No reutilizable','Pendiente de revisión'])
            reintegrar=st.checkbox('Reintegrar unidades al inventario disponible',value=False)
            observaciones=st.text_area('Observaciones')
            guardar=st.form_submit_button('Registrar devolución')
        if guardar:
            if reintegrar and condicion!='Reutilizable': st.error('Solo las prendas reutilizables pueden reintegrarse al inventario.')
            else:
                con=conectar()
                try:
                    con.execute('BEGIN IMMEDIATE')
                    con.execute("""INSERT INTO devoluciones (empleado_id,prenda_id,cantidad,fecha,motivo,condicion,reintegrada,usuario,observaciones)
                        VALUES (?,?,?,?,?,?,?,?,?)""",(empleado[0],prenda[0],cantidad,ahora(),motivo,condicion,int(reintegrar),st.session_state.usuario,observaciones.strip()))
                    if reintegrar:
                        con.execute('UPDATE prendas SET existencia=existencia+? WHERE id=?',(cantidad,prenda[0]))
                        con.execute("""INSERT INTO movimientos (prenda_id,tipo,cantidad,fecha,usuario,descripcion)
                            VALUES (?,'ENTRADA',?,?,?,?)""",(prenda[0],cantidad,ahora(),st.session_state.usuario,f'Devolución reutilizable: {motivo}'))
                    con.commit(); registrar_bitacora('Devoluciones','Registro',f'{empleado[1]} | {prenda[1]} | {cantidad} | Reintegrada: {reintegrar}')
                    st.success('Devolución registrada correctamente.'); st.rerun()
                except Exception as exc: con.rollback(); st.error(str(exc))
                finally: con.close()
    con=conectar()
    historial=con.execute("""SELECT d.id,d.fecha,e.numero_empleado,e.nombre,p.nombre,p.talla,d.cantidad,d.motivo,d.condicion,d.reintegrada,d.usuario
        FROM devoluciones d JOIN empleados e ON e.id=d.empleado_id JOIN prendas p ON p.id=d.prenda_id ORDER BY d.id DESC""").fetchall()
    con.close()
    if historial:
        import pandas as pd
        st.dataframe(pd.DataFrame(historial,columns=['Folio','Fecha','Matrícula','Empleado','Prenda','Talla','Cantidad','Motivo','Condición','Reintegrada','Usuario']),hide_index=True,use_container_width=True)


# ============================================================
# REPORTES
# ============================================================
elif opcion == "📑 Reportes":

    st.title("📑 Reportes del sistema")
    st.write("Genera consultas descargables de inventario, entregas y movimientos.")

    import pandas as pd
    from datetime import date

    fecha_inicio = st.date_input("Fecha inicial", value=date.today().replace(day=1), key="rep_inicio")
    fecha_fin = st.date_input("Fecha final", value=date.today(), key="rep_fin")

    if fecha_inicio > fecha_fin:
        st.error("❌ La fecha inicial no puede ser posterior a la fecha final.")
    else:
        inicio_txt = str(fecha_inicio)
        fin_txt = str(fecha_fin) + " 23:59:59"

        conexion = conectar()
        cursor = conexion.cursor()

        cursor.execute("""
            SELECT e.numero_empleado AS Matricula, e.nombre AS Empleado,
                   e.area AS Area, e.puesto AS Puesto,
                   p.nombre AS Prenda, p.talla AS Talla,
                   en.cantidad AS Cantidad, en.fecha AS Fecha,
                   en.responsable AS Responsable
            FROM entregas en
            INNER JOIN empleados e ON en.empleado_id = e.id
            INNER JOIN prendas p ON en.prenda_id = p.id
            WHERE datetime(en.fecha) BETWEEN datetime(?) AND datetime(?)
            ORDER BY datetime(en.fecha) DESC
        """, (inicio_txt, fin_txt))
        entregas_rep = cursor.fetchall()

        cursor.execute("""
            SELECT m.id AS Folio, m.fecha AS Fecha, m.tipo AS Movimiento,
                   p.nombre AS Prenda, p.talla AS Talla, m.cantidad AS Cantidad,
                   m.usuario AS Usuario, m.descripcion AS Descripcion
            FROM movimientos m
            INNER JOIN prendas p ON m.prenda_id = p.id
            WHERE datetime(m.fecha) BETWEEN datetime(?) AND datetime(?)
            ORDER BY datetime(m.fecha) DESC
        """, (inicio_txt, fin_txt))
        movimientos_rep = cursor.fetchall()

        cursor.execute("""
            SELECT nombre AS Prenda, tipo AS Tipo, talla AS Talla,
                   existencia AS Existencia, minimo AS Minimo
            FROM prendas
            WHERE activo = 1
            ORDER BY nombre, talla
        """)
        inventario_rep = cursor.fetchall()
        conexion.close()

        df_ent = pd.DataFrame(entregas_rep, columns=["Matrícula","Empleado","Área","Puesto","Prenda","Talla","Cantidad","Fecha","Responsable"])
        df_mov = pd.DataFrame(movimientos_rep, columns=["Folio","Fecha","Movimiento","Prenda","Talla","Cantidad","Usuario","Descripción"])
        df_inv = pd.DataFrame(inventario_rep, columns=["Prenda","Tipo","Talla","Existencia","Mínimo"])

        t1, t2, t3 = st.tabs(["📦 Entregas", "🔄 Movimientos", "👕 Inventario"])
        with t1:
            st.metric("Entregas en periodo", len(df_ent))
            st.dataframe(df_ent, use_container_width=True, hide_index=True)
            if not df_ent.empty:
                st.download_button("📥 Descargar entregas CSV", df_ent.to_csv(index=False).encode("utf-8-sig"), "reporte_entregas.csv", "text/csv", use_container_width=True)
        with t2:
            st.metric("Movimientos en periodo", len(df_mov))
            st.dataframe(df_mov, use_container_width=True, hide_index=True)
            if not df_mov.empty:
                st.download_button("📥 Descargar movimientos CSV", df_mov.to_csv(index=False).encode("utf-8-sig"), "reporte_movimientos.csv", "text/csv", use_container_width=True)
        with t3:
            st.metric("Prendas activas", len(df_inv))
            st.dataframe(df_inv, use_container_width=True, hide_index=True)
            if not df_inv.empty:
                st.download_button("📥 Descargar inventario CSV", df_inv.to_csv(index=False).encode("utf-8-sig"), "reporte_inventario.csv", "text/csv", use_container_width=True)

        st.divider()
        st.subheader('📥 Reporte consolidado en Excel y PDF')
        hojas={'Entregas':df_ent,'Movimientos':df_mov,'Inventario':df_inv}
        try:
            excel=generar_excel_reportes(hojas)
            st.download_button('📗 Descargar Excel completo',excel,f'TNG_reporte_{fecha_inicio}_{fecha_fin}.xlsx',
                mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',use_container_width=True)
        except ImportError: st.warning('Para Excel instala: pip install openpyxl')
        try:
            pdf=generar_pdf_reporte(hojas,fecha_inicio,fecha_fin)
            st.download_button('📕 Descargar resumen PDF',pdf,f'TNG_reporte_{fecha_inicio}_{fecha_fin}.pdf',
                mime='application/pdf',use_container_width=True)
        except ImportError: st.warning('Para PDF instala: pip install reportlab')

        st.divider()
        st.subheader("📊 Resumen del periodo")
        r1, r2, r3, r4 = st.columns(4)
        r1.metric("Entregas", len(df_ent))
        r2.metric("Unidades entregadas", int(df_ent["Cantidad"].sum()) if not df_ent.empty else 0)
        r3.metric("Movimientos", len(df_mov))
        r4.metric("Prendas bajo mínimo", int((df_inv["Existencia"] <= df_inv["Mínimo"]).sum()) if not df_inv.empty else 0)


# ============================================================
# NUEVOS MÓDULOS (migraciones aditivas)
# ============================================================
elif opcion == "🗂️ Expediente y bajas":
    st.title("🗂️ Expediente digital y control de bajas")
    with conectar() as con:
        empleados=con.execute("SELECT id,numero_empleado,nombre,puesto,area,disciplina,tipo_trabajador,talla_camisa,talla_pantalon,fecha_ingreso,activo FROM empleados ORDER BY nombre").fetchall()
    if not empleados: st.info("No hay trabajadores registrados.")
    else:
        filtrado=st.text_input("Buscar por nombre o matrícula",key="exp_buscar").casefold().strip()
        candidatos=[e for e in empleados if filtrado in (e[1]+' '+e[2]).casefold()]
        if not candidatos: st.info("Sin coincidencias.")
        else:
            elegido=st.selectbox("Trabajador",candidatos,format_func=lambda e:f"{e[1]} — {e[2]} ({'Activo' if e[-1] else 'Inactivo'})")
            eid,mat,nombre,puesto,area,disc,tipo,cam,pant,ingreso,activo=elegido
            st.write({"Matrícula":mat,"Nombre":nombre,"Puesto":puesto,"Área":area,"Disciplina":disc,"Tipo":tipo,"Camisola":cam,"Pantalón":pant,"Ingreso":ingreso,"Estado":"Activo" if activo else "Inactivo"})
            with conectar() as con:
                for titulo,sql in [
                    ('Entregas',"SELECT e.fecha,p.nombre,p.talla,e.cantidad,e.responsable FROM entregas e JOIN prendas p ON p.id=e.prenda_id WHERE e.empleado_id=? ORDER BY e.fecha DESC"),
                    ('Solicitudes',"SELECT s.fecha_solicitud,p.nombre,p.talla,s.cantidad,s.estado FROM solicitudes_almacen s JOIN prendas p ON p.id=s.prenda_id WHERE s.empleado_id=? ORDER BY s.fecha_solicitud DESC"),
                    ('Reposiciones',"SELECT r.fecha_solicitud,p.nombre,p.talla,r.cantidad,r.motivo,r.estado FROM reposiciones r JOIN prendas p ON p.id=r.prenda_id WHERE r.empleado_id=? ORDER BY r.fecha_solicitud DESC"),
                    ('Devoluciones',"SELECT d.fecha,p.nombre,p.talla,d.cantidad,d.motivo FROM devoluciones d JOIN prendas p ON p.id=d.prenda_id WHERE d.empleado_id=? ORDER BY d.fecha DESC"),
                    ('Cambios de talla',"SELECT fecha,campo,anterior,nueva,motivo,usuario FROM cambios_talla WHERE empleado_id=? ORDER BY fecha DESC")]:
                    with st.expander(titulo):
                        st.dataframe(con.execute(sql,(eid,)).fetchall(),use_container_width=True)
            if activo:
                with st.form('baja_form'):
                    motivo_baja=st.text_area('Motivo de baja *')
                    confirmar=st.checkbox('Confirmo la baja del trabajador')
                    if st.form_submit_button('Registrar baja'):
                        if not motivo_baja.strip() or not confirmar: st.error('Indica el motivo y confirma la baja.')
                        else:
                            with conectar() as con:
                                con.execute('UPDATE empleados SET activo=0 WHERE id=? AND activo=1',(eid,))
                                con.execute('INSERT INTO bajas_trabajadores(empleado_id,fecha,motivo,usuario) VALUES(?,?,?,?)',(eid,ahora(),motivo_baja.strip(),st.session_state.usuario))
                            registrar_bitacora('Empleados','Baja',f'Matrícula {mat}: {motivo_baja.strip()}')
                            st.success('Trabajador desactivado; se conserva su historial.');st.rerun()
            else: st.info('Trabajador inactivo. Su expediente histórico permanece disponible.')

elif opcion == "📏 Cambios de talla":
    st.title('📏 Historial de cambios de talla')
    with conectar() as con:
        trabajadores=con.execute('SELECT id,numero_empleado,nombre,talla_camisa,talla_pantalon FROM empleados WHERE activo=1 ORDER BY nombre').fetchall()
    if trabajadores:
        with st.form('cambio_talla'):
            e=st.selectbox('Trabajador',trabajadores,format_func=lambda r:f'{r[1]} — {r[2]}')
            campo=st.selectbox('Prenda',['talla_camisa','talla_pantalon'],format_func=lambda c:'Camisola' if c=='talla_camisa' else 'Pantalón')
            nueva=st.selectbox('Nueva talla',[str(n) for n in range(28,53,2)])
            motivo=st.text_input('Motivo *')
            if st.form_submit_button('Guardar cambio'):
                anterior=e[3] if campo=='talla_camisa' else e[4]
                if not motivo.strip(): st.error('El motivo es obligatorio.')
                elif anterior==nueva: st.info('La talla seleccionada es igual a la actual.')
                else:
                    with conectar() as con:
                        con.execute(f'UPDATE empleados SET {campo}=? WHERE id=?',(nueva,e[0]))
                        con.execute('INSERT INTO cambios_talla(empleado_id,campo,anterior,nueva,motivo,fecha,usuario) VALUES(?,?,?,?,?,?,?)',(e[0],campo,anterior,nueva,motivo.strip(),ahora(),st.session_state.usuario))
                    registrar_bitacora('Empleados','Cambio de talla',f'{e[1]} {campo}: {anterior} -> {nueva}')
                    st.success('Cambio guardado.');st.rerun()
    else: st.info('No hay trabajadores activos.')

elif opcion == "⚙️ Motivos de reposición":
    st.title('⚙️ Catálogo de motivos de reposición')
    with conectar() as con:
        filas=con.execute('SELECT id,motivo,activo FROM catalogo_motivos_reposicion ORDER BY motivo').fetchall()
    st.dataframe([{'ID':i,'Motivo':m,'Activo':bool(a)} for i,m,a in filas],hide_index=True,use_container_width=True)
    with st.form('nuevo_motivo'):
        nombre=st.text_input('Nuevo motivo *')
        if st.form_submit_button('Agregar motivo'):
            if not nombre.strip(): st.error('Escribe un motivo.')
            else:
                try:
                    with conectar() as con: con.execute('INSERT INTO catalogo_motivos_reposicion(motivo) VALUES(?)',(nombre.strip(),))
                    registrar_bitacora('Reposiciones','Nuevo motivo',nombre.strip());st.rerun()
                except sqlite3.IntegrityError: st.error('Ese motivo ya existe.')
    if filas:
        seleccionado=st.selectbox('Activar/desactivar motivo',filas,format_func=lambda r:f'{r[1]} — {"Activo" if r[2] else "Inactivo"}')
        if st.button('Cambiar estado'):
            with conectar() as con: con.execute('UPDATE catalogo_motivos_reposicion SET activo=? WHERE id=?',(0 if seleccionado[2] else 1,seleccionado[0]))
            registrar_bitacora('Reposiciones','Estado motivo',seleccionado[1]);st.rerun()

elif opcion == "💾 Respaldos":
    st.title('💾 Respaldo de la base de datos')
    st.info('Se genera un respaldo diario en la carpeta respaldos_tng, junto a la base de datos. Mantén también copias en una ubicación segura distinta.')
    if st.button('Crear respaldo ahora'):
        destino=Path(DB_NAME).resolve().parent/'respaldos_tng';destino.mkdir(exist_ok=True)
        ruta=destino/('tng_uniformes_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.db')
        with conectar() as origen, sqlite3.connect(str(ruta)) as copia: origen.backup(copia)
        st.success(f'Respaldo creado: {ruta.name}')
        registrar_bitacora('Respaldos','Respaldo manual',ruta.name)
    ruta=respaldo_automatico()
    st.caption(f'Respaldo diario: {ruta.name}')
    st.download_button('Descargar respaldo diario',data=ruta.read_bytes(),file_name=ruta.name,mime='application/octet-stream')

elif opcion == "🔔 Notificaciones":
    st.title('🔔 Centro de notificaciones')
    with conectar() as con:
        bajas=con.execute('SELECT nombre,talla,existencia,minimo FROM prendas WHERE activo=1 AND existencia<=minimo ORDER BY existencia').fetchall()
        solicitudes=con.execute("SELECT estado,COUNT(*) FROM solicitudes_almacen WHERE estado NOT IN ('Entregada','Rechazada','Cancelada') GROUP BY estado").fetchall()
        proximas=calcular_proximas_dotaciones(con)
    st.subheader('Existencias bajas')
    if bajas: st.dataframe([{'Prenda':n,'Talla':t,'Existencia':e,'Mínimo':m} for n,t,e,m in bajas],hide_index=True,use_container_width=True)
    else: st.success('Sin alertas de inventario bajo.')
    st.subheader('Solicitudes abiertas')
    st.write(dict(solicitudes) if solicitudes else 'No hay solicitudes abiertas.')
    st.subheader('Dotaciones vencidas o próximas en 30 días')
    limite=date.today()+timedelta(days=30)
    alertas=[r for r in proximas if r['Próxima fecha']!='Por definir' and r['Próxima fecha']<=limite.isoformat()]
    st.dataframe(alertas,use_container_width=True,hide_index=True) if alertas else st.success('Sin dotaciones próximas identificadas.')

elif opcion == "📊 Indicadores KPI":
    st.title('📊 Indicadores de gestión')
    with conectar() as con:
        activos=con.execute('SELECT COUNT(*) FROM empleados WHERE activo=1').fetchone()[0]
        prendas=con.execute('SELECT COUNT(*) FROM prendas WHERE activo=1').fetchone()[0]
        bajas=con.execute('SELECT COUNT(*) FROM prendas WHERE activo=1 AND existencia<=minimo').fetchone()[0]
        pendientes=con.execute("SELECT COUNT(*) FROM solicitudes_almacen WHERE estado NOT IN ('Entregada','Rechazada','Cancelada')").fetchone()[0]
        entregadas=con.execute('SELECT COALESCE(SUM(cantidad),0) FROM entregas').fetchone()[0]
        total=con.execute('SELECT COUNT(*) FROM solicitudes_almacen').fetchone()[0]
        completas=con.execute("SELECT COUNT(*) FROM solicitudes_almacen WHERE estado='Entregada'").fetchone()[0]
    c1,c2,c3=st.columns(3)
    c1.metric('Trabajadores activos',activos);c2.metric('Prendas en catálogo',prendas);c3.metric('Prendas con alerta',bajas)
    c4,c5,c6=st.columns(3)
    c4.metric('Solicitudes abiertas',pendientes);c5.metric('Unidades entregadas (histórico)',entregadas);c6.metric('Solicitudes completadas',f'{completas}/{total}')
    st.metric('Porcentaje de solicitudes completadas',f'{(100*completas/total):.1f}%' if total else 'Sin datos')
    st.caption('Los KPI se calculan a partir de los registros de la base de datos; no son valores simulados.')

elif opcion == "🤖 Asistente TNG 2.0":
    st.title('🤖 Asistente TNG')
    st.caption('Asistente local de consulta. No usa IA externa, no modifica registros y no inventa matrículas.')
    if 'tng_chat' not in st.session_state: st.session_state.tng_chat=[]
    for pregunta,respuesta in st.session_state.tng_chat:
        with st.chat_message('user'): st.write(pregunta)
        with st.chat_message('assistant'): st.markdown(respuesta)
    pregunta=st.chat_input('Ejemplo: Busca Juan Pérez / ¿Hay botas DURAMAX 27?')
    if pregunta:
        respuesta=buscar_asistente(pregunta,rol)
        st.session_state.tng_chat.append((pregunta,respuesta))
        with st.chat_message('user'): st.write(pregunta)
        with st.chat_message('assistant'): st.markdown(respuesta)

# ============================================================
# BITÁCORA
# ============================================================
elif opcion == "📝 Bitácora":

    st.title("📝 Bitácora de acciones")
    st.write("Historial de acciones importantes realizadas dentro del sistema.")

    conexion = conectar()
    cursor = conexion.cursor()
    cursor.execute("""
        SELECT id, fecha, usuario, rol, modulo, accion, detalle
        FROM bitacora
        ORDER BY datetime(fecha) DESC, id DESC
        LIMIT 1000
    """)
    bitacora_data = cursor.fetchall()
    conexion.close()

    import pandas as pd
    df_b = pd.DataFrame(bitacora_data, columns=["Folio","Fecha","Usuario","Rol","Módulo","Acción","Detalle"])

    if df_b.empty:
        st.info("Todavía no hay acciones registradas en la bitácora.")
    else:
        usuarios = ["Todos"] + sorted(df_b["Usuario"].dropna().unique().tolist())
        modulos = ["Todos"] + sorted(df_b["Módulo"].dropna().unique().tolist())
        c1, c2 = st.columns(2)
        with c1:
            filtro_usuario = st.selectbox("Usuario", usuarios, key="bit_usuario")
        with c2:
            filtro_modulo = st.selectbox("Módulo", modulos, key="bit_modulo")

        vista = df_b.copy()
        if filtro_usuario != "Todos":
            vista = vista[vista["Usuario"] == filtro_usuario]
        if filtro_modulo != "Todos":
            vista = vista[vista["Módulo"] == filtro_modulo]

        st.metric("Acciones encontradas", len(vista))
        st.dataframe(vista, use_container_width=True, hide_index=True)
        st.download_button("📥 Descargar bitácora CSV", vista.to_csv(index=False).encode("utf-8-sig"), "bitacora.csv", "text/csv", use_container_width=True)
